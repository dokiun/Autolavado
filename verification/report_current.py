"""Workbook by equipment, from the current OpenModelica run only."""
import csv
import hashlib
import math
import textwrap
import xml.etree.ElementTree as ET
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'Memoria_Calculo_Hidraulica_Autolavado.xlsx'
RESULT = ROOT / 'verification/current/CurrentCycle_res.csv'
XML = RESULT.with_name('CurrentCycle_init.xml')
OUTPUT = ROOT / 'Memoria_Calculo_Hidraulica_Modelica.xlsx'
before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
variables = {v.attrib['name']: v for v in ET.parse(XML).iter('ScalarVariable')}
targets = [9, 18, 33, 48, 55, 58, 64]
samples = {}
distances = {t: math.inf for t in targets}
end = 0
minimum_nozzle_pressure = math.inf
maximum_mass_residual = 0
maximum_energy_residual = 0
with RESULT.open(newline='') as f:
    for raw in csv.DictReader(f):
        t = float(raw['time'])
        end = max(end, t)
        minimum_nozzle_pressure = min(minimum_nozzle_pressure,float(raw['plant.gun.nozzle.port_a.p']))
        maximum_mass_residual = max(maximum_mass_residual,abs(float(raw['plant.mainMassResidual'])),abs(float(raw['plant.chemicalMassResidual'])))
        maximum_energy_residual = max(maximum_energy_residual,abs(float(raw['plant.inletHeater.energyResidual'])))
        for target in targets:
            if abs(t-target) < distances[target]:
                distances[target] = abs(t-target)
                samples[target] = {k: float(v) for k,v in raw.items()}
assert end >= 64 and all(d < .021 for d in distances.values())
assert all(math.isfinite(v) for row in samples.values() for v in row.values())
assert minimum_nozzle_pressure>0,'Nonphysical nozzle pressure'
assert maximum_mass_residual<1e-6 and maximum_energy_residual<1e-4,'Conservation residuals'

def get(k, time=9):
    if k in samples[time]:
        return samples[time][k]
    v = variables.get(k)
    if v is None:
        return None
    if v.attrib.get('alias') in ('alias','negatedAlias'):
        result = get(v.attrib['aliasVariable'], time)
        return result if result is None else result * (-1 if v.attrib['alias']=='negatedAlias' else 1)
    child = next(iter(v), None)
    if child is not None and v.attrib.get('variability') in ('parameter','constant'):
        raw = child.attrib.get('start')
        if raw is not None:
            try:
                return float(raw)
            except ValueError:
                return {'true':True,'false':False}.get(raw)
    if k == 'plant.salida.p':
        return get('plant.pAtmosphere',time)
    if k == 'plant.gun.nozzleArea':
        return math.pi*get('plant.lowPressureNozzleDiameter',time)**2/4
    if k.endswith('.siphonLength'):
        return get('plant.chemicalSiphonLength',time)
    if k.endswith('.siphonDiameter'):
        return get('plant.chemicalSiphonDiameter',time)
    return None

assert abs(get('plant.nozzleDiameter')-.00115) < 1e-12, 'Wrong nozzle in simulation'
template = openpyxl.load_workbook(SOURCE)['Bomba']
book = openpyxl.Workbook()
book.remove(book.active)
rho = get('plant.gun.nozzle.rho_nominal')
assert rho is not None
patm = get('plant.pAtmosphere')
assert patm is not None

def data(label, key, unit, factor=1, offset=0, time=9, note=''):
    value = get(key, time)
    v = variables.get(key)
    is_parameter = v is not None and v.attrib.get('variability') in ('parameter','constant')
    state = 'Parámetro del modelo' if is_parameter else 'Resultado simulado'
    source = key + ('' if is_parameter else f'; t = {samples[time]["time"]:g} s')
    if note:
        source += '. ' + note
    if value is None:
        state = 'No disponible'
        source = key + '; sin valor exportado'
    else:
        value = value*factor+offset
    return (label,value,unit,state,source)

def pressure(label,key,time=9):
    return data(label,key,'bar(g)',1e-5,-patm/1e5,time)

def text(label,value,note):
    return (label,value,'—','Descripción del modelo',note)

def copy_cell_style(destination, source):
    # Register each component in the destination workbook; style IDs are local.
    for attribute in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(destination, attribute, copy(getattr(source, attribute)))
    destination.number_format = source.number_format

def sheet(name,title,rows,time=9):
    s = book.create_sheet(name)
    s.sheet_view.showGridLines = False
    s.merge_cells('A1:E1')
    s['A1'] = title
    copy_cell_style(s['A1'], template['A1'])
    s.merge_cells('A2:E2')
    s['A2'] = f'Modelo actual · Boquillas: alta 1,15 mm / baja 4,6 mm · Simulación CurrentCycle · Caso de esta hoja: t = {time} s'
    copy_cell_style(s['A2'], template['A2'])
    s['A2'].alignment = Alignment(wrap_text=True,vertical='center')
    s.row_dimensions[1].height = 32
    s.row_dimensions[2].height = 36
    for col,label in enumerate(['Parámetro','Valor','Unidad','Estado','Fuente / observación'],1):
        c = s.cell(4,col,label)
        copy_cell_style(c, template.cell(4,col))
        c.alignment = Alignment(wrap_text=True,vertical='center',horizontal='center')
    s.row_dimensions[4].height = 30
    widths = [40,24,16,27,68]
    for col,width in zip('ABCDE',widths):
        s.column_dimensions[col].width = width
    for r,values in enumerate(rows,5):
        lines = 1
        for col,value in enumerate(values,1):
            c = s.cell(r,col,value)
            copy_cell_style(c, template.cell(5,col))
            c.alignment = Alignment(vertical='center',wrap_text=True,horizontal='right' if col==2 and isinstance(value,(int,float)) else 'left')
            if isinstance(value,(int,float)) and not isinstance(value,bool):
                c.number_format = '#,##0.000;[Red]-#,##0.000;0.000' if value and abs(value)<100 else '#,##0.00;[Red]-#,##0.00;0.00'
            lines = max(lines, len(textwrap.wrap(str(value or ''),max(8,int(widths[col-1]*.9)))))
        s.row_dimensions[r].height = max(30,lines*16+10)
    s.freeze_panes = 'B5'
    s.auto_filter.ref = f'A4:E{s.max_row}'
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_setup.orientation = 'landscape'
    s.page_setup.paperSize = s.PAPERSIZE_A4
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.print_title_rows = '1:4'
    s.print_options.horizontalCentered = True
    s.print_area = f'A1:E{s.max_row}'
    return s

sheet('Resumen','RESULTADOS DEL MODELO ACTUAL',[
    text('Caso principal','Enjuague','Instantánea del ciclo dinámico a t = 9 s; no se declara régimen estacionario.'),
    data('Diámetro de la boquilla','plant.nozzleDiameter','mm',1000),
    data('Caudal de agua','plant.waterFlow_L_min','L/min'),
    data('Caudal de descarga','plant.outletFlow_L_min','L/min'),
    data('Caudal químico al inyector','plant.injector.mChemical','L/min',60000/rho),
    data('Aspiracion durante prelavado','plant.injector.mChemical','L/min',60000/rho,time=18),
    data('Aspiracion durante jabon','plant.injector.mChemical','L/min',60000/rho,time=33),
    data('Aspiracion durante cera','plant.injector.mChemical','L/min',60000/rho,time=48),
    pressure('Presión de descarga de bomba','plant.pumpOutletPressure'),
    pressure('Presión en garganta del Venturi','plant.throatPressure'),
    data('Temperatura de entrada a bomba','plant.inletHeater.T','°C',1,-273.15),
    data('Temperatura de descarga','plant.outletTemperature','°C',1,-273.15),
    data('Temperatura final del calentador (sin flujo)','plant.inletHeater.T','°C',1,-273.15,time=64,note='Agua retenida: pistola cerrada y caudal nulo.'),
    data('Temperatura del calentador al final con caudal','plant.inletHeater.T','°C',1,-273.15,time=55,note='Último instante a plena apertura, antes del cierre de la pistola.'),
])
sheet('Alimentacion','ALIMENTACIÓN DE AGUA',[
    text('Equipo','Fuente de agua','Frontera hidráulica y térmica del modelo.'),
    data('Temperatura de suministro','plant.supplyTemperature','°C',1,-273.15),
    data('Presión absoluta de suministro','plant.pAtmosphere','kPa(abs)',.001),
    data('Caudal de agua','plant.waterFlow_L_min','L/min'),
    text('Medio','Agua de propiedades constantes','Modelica.Media.Water.ConstantPropertyLiquidWater'),
    data('Densidad','plant.gun.nozzle.rho_nominal','kg/m³'),
])
sheet('Calentador','CALENTADOR EXTERNO HACEB',[
    text('Equipo','HACEB · Paso a gas','HacebGasHeater; referencia exacta no identificada.'),
    data('Ajuste manual de gas','plant.inletHeater.gasSetting','0–1',time=64),
    data('Caudal nominal','plant.inletHeater.ratedFlow','L/min',60000,time=64),
    data('Potencia de gas aplicada al final','plant.inletHeater.gasPower','kW',.001,time=64),
    data('Potencia térmica del quemador al final','plant.inletHeater.heaterPower','kW',.001,time=64),
    data('Potencia transferida al agua al final','plant.inletHeater.heatToWater','kW',.001,time=64),
    data('Temperatura final del agua retenida','plant.inletHeater.T','°C',1,-273.15,time=64,note='Final de simulación; sin circulación de agua.'),
    data('Caudal másico al final','plant.inletHeater.m_flow','kg/s',time=64),
    pressure('Presión de salida al final','plant.pumpInletPressure',time=64),
    data('Temperatura al final con caudal','plant.inletHeater.T','°C',1,-273.15,time=55,note='Último instante a plena apertura antes del cierre; no es la temperatura del agua retenida.'),
    data('Caudal másico antes del cierre','plant.inletHeater.m_flow','kg/s',time=55),
],64)
sheet('Hidrolavadora','HIDROLAVADORA TRENT HLT395-1800',[
    text('Equipo','TRENT HLT395-1800','PressureWasherPump: modelo equivalente de la hidrolavadora completa.'),
    data('Potencia eléctrica nominal','plant.bomba.ratedElectricalPower','W'),
    data('Caudal máximo parametrizado','plant.bomba.maximumFlow','L/min',60000),
    data('Rendimiento hidráulico parametrizado','plant.bomba.hydraulicEfficiency','0–1'),
    data('Presión máxima parametrizada','plant.maximumPressure','bar',1e-5),
    pressure('Presión de aspiración','plant.pumpInletPressure'),
    pressure('Presión de descarga','plant.pumpOutletPressure'),
    data('Caudal de agua','plant.waterFlow_L_min','L/min'),
    data('Temperatura de aspiración','plant.inletHeater.T','°C',1,-273.15),
    ( 'NPSHr',None,'m','No modelado','El modelo no exporta NPSHr; no se asigna un valor supuesto.'),
])
sheet('Venturi','INYECTOR VENTURI',[
    text('Equipo','General Pump 100819','Referencia downstream elegida: https://www.generalpump.com/product/100819/. Equivalente hidraulico pendiente de calibracion; no hay curva de aspiracion publicada en la ficha.'),
    data('Orificio interno del inyector','plant.injector.commercialOrificeDiameter','mm',1000,note='No es la boquilla de descarga de 1,15 mm.'),
    data('Caudal minimo de catalogo','plant.injector.commercialMinimumFlow','L/min',60000),
    data('Caudal maximo de catalogo','plant.injector.commercialMaximumFlow','L/min',60000),
    data('Rating de presion del cuerpo','plant.injector.commercialPressureRating','bar',1e-5,note='No equivale a contrapresion admisible para aspirar.'),
    pressure('Presión de entrada','plant.injectorInletPressure'),
    pressure('Presión de salida','plant.injectorOutletPressure'),
    pressure('Presión de garganta / aspiración','plant.throatPressure'),
    pressure('Presión del manifold químico','plant.chemicalManifoldPressure'),
    data('Pérdida de presión','plant.injectorPressureDrop','kPa',.001),
    data('Caudal de agua','plant.waterFlow_L_min','L/min'),
    data('Caudal de químico aspirado','plant.injector.mChemical','L/min',60000/rho),
    data('Caudal de mezcla descargada','plant.outletFlow_L_min','L/min'),
    pressure('Presion de salida durante jabon','plant.injectorOutletPressure',33),
    pressure('Presion de garganta durante jabon','plant.throatPressure',33),
    data('Aspiracion durante jabon','plant.injector.mChemical','L/min',60000/rho,time=33),
])
for title,prefix,time in [('Prelavado','prewash',18),('Jabon','soap',33),('Cera','wax',48)]:
    p='plant.'+prefix
    sheet(title,'SISTEMA DE INYECCIÓN · '+title.upper(),[
        text('Equipo','Tanque y rama '+title,'Tanque con respiradero > sifón > electroválvula NC > check > aguja > manifold.'),
        data('Volumen inicial del tanque',p+'.tankLiquidVolume','L',1000,time=time),
        data('Sección del tanque',p+'.tankCrossSection','m²',time=time),
        data('Longitud del sifón',p+'.siphonLength','m',time=time),
        data('Diámetro interior del sifón',p+'.siphonDiameter','mm',1000,time=time),
        data('Ajuste manual de aguja',p+'.needleSetting','0–1',time=time),
        data('Apertura de electroválvula',p+'.solenoid.opening','0–1',time=time),
        pressure('Presión antes de electroválvula',p+'.valveInletPressure',time),
        pressure('Presión después de electroválvula',p+'.valveOutletPressure',time),
        data('Diferencial en electroválvula',p+'.valveDifferentialPressure','bar',1e-5,time=time),
        data('Caudal extraído del tanque',p+'.tankDraw','L/min',60000/rho,time=time),
        data('Caudal en válvula check',p+'.checkValve.m_flow','L/min',60000/rho,time=time),
        data('Caudal en aguja',p+'.needle.m_flow','L/min',60000/rho,time=time),
        data('Inyeccion bloqueada por diferencial','plant.chemicalInjectionBlocked','0=no; 1=si',time=time,note='Diferencia insuficiente entre manifold y garganta, aunque la electrovalvula abra.'),
    ],time)
sheet('Manguera','MANGUERA DE ALTA PRESIÓN',[
    text('Equipo','StaticPipe','plant.hose; tubería equivalente de la manguera.'),
    data('Longitud','plant.hoseLength','m'),
    data('Diámetro interior','plant.hoseDiameter','mm',1000),
    data('Rugosidad','plant.hose.roughness','mm',1000),
    data('Caudal','plant.outletFlow_L_min','L/min'),
    data('Velocidad media','plant.hose.flowModel.vs[1]','m/s'),
    data('Pérdida de presión','plant.hosePressureDrop','kPa',.001),
    pressure('Presión de entrada','plant.injectorOutletPressure'),
    pressure('Presión de salida','plant.gunInletPressure'),
])
sheet('Pistola','PISTOLA DE DESCARGA',[
    text('Equipo','ValveIncompressible','plant.gun.pistola; apertura comandada por el gatillo.'),
    data('Apertura','plant.gun.pistola.opening','0–1'),
    pressure('Presión de entrada','plant.gunInletPressure'),
    pressure('Presión de salida','plant.gun.nozzle.port_a.p'),
    data('Pérdida de presión','plant.gun.pistola.dp','kPa',.001),
    data('Caudal de descarga','plant.outletFlow_L_min','L/min'),
])
sheet('Boquilla','BOQUILLAS DE ALTA Y BAJA PRESIÓN',[
    text('Equipo','Seleccion de dos boquillas','Alta para enjuague; baja para prelavado, jabon y cera. Seleccion ideal por programa; el modelo supone que el operador instala la boquilla correspondiente.'),
    data('Diámetro de orificio','plant.nozzleDiameter','mm',1000),
    data('Diametro de boquilla de baja','plant.lowPressureNozzleDiameter','mm',1000,note='Supuesto de diseno: confirmar con boquilla comercial.'),
    data('Coeficiente de descarga','plant.nozzleDischargeCoefficient','—',note='Supuesto del modelo.'),
    data('Área activa en enjuague','plant.gun.activeNozzleArea','mm²',1e6),
    pressure('Presión de entrada','plant.gun.nozzle.port_a.p'),
    pressure('Presión de salida','plant.gun.nozzle.port_b.p'),
    data('Diferencial de presión','plant.nozzlePressureDrop','bar',1e-5),
    data('Caudal de descarga','plant.outletFlow_L_min','L/min'),
    data('Temperatura de descarga','plant.outletTemperature','°C',1,-273.15),
    data('Diametro activo en aplicacion','plant.gun.activeNozzleDiameter','mm',1000,time=33),
    data('Presion en boquilla de baja','plant.nozzlePressureDrop','bar',1e-5,time=33),
    data('Caudal descargado en aplicacion','plant.outletFlow_L_min','L/min',time=33),
])
sheet('Pistola_cerrada','CIRCUITO CON PISTOLA CERRADA',[
    data('Apertura de pistola','plant.gun.pistola.opening','0–1',time=58),
    pressure('Presión de descarga de bomba','plant.pumpOutletPressure',58),
    pressure('Presión de garganta','plant.throatPressure',58),
    data('Caudal de agua','plant.waterFlow_L_min','L/min',time=58),
    data('Caudal de descarga','plant.outletFlow_L_min','L/min',time=58),
    data('Caudal por bypass','plant.bypassFlow_L_min','L/min',time=58),
],58)
def number(value, digits=3):
    return f'{value:.{digits}f}'.replace('.', ',')

def pg(key, time=9):
    return (get(key,time)-patm)/1e5

def temp(key, time=9):
    return get(key,time)-273.15

def scaled(key, factor, time=9):
    return get(key,time)*factor

decisions = sheet('Decisiones_diseno', 'DECISIONES DE DISEÑO BASADAS EN LA SIMULACIÓN', [
    ('Boquillas', 'Usar alta de 1,15 mm en enjuague y baja de 4,6 mm al aplicar químicos.',
     'Alta: '+number(pg('plant.injectorInletPressure'))+' bar(g) en entrada del Venturi a 9 s. Baja: '+number(pg('plant.injectorOutletPressure',33))+' bar(g) en salida a 33 s.',
     'Decisión de diseño provisional', 'plant.gun.activeNozzleDiameter; plant.injectorInletPressure / OutletPressure. Confirmar diámetro de baja y Cd con la boquilla comercial.'),
    ('Inyección downstream', 'Conservar el inyector después de la hidrolavadora y aplicar químicos con baja contrapresión.',
     'Garganta en jabón: '+number(pg('plant.throatPressure',33))+' bar(g); aspiración: '+number(scaled('plant.injector.mChemical',60000/rho,time=33))+' L/min.',
     'Respaldado por el modelo', 'plant.throatPressure y plant.injector.mChemical, t=33 s. General Pump 100819 es referencia comercial; los coeficientes hidráulicos siguen sin calibrar.'),
    ('Dosificación', 'Mantener ajuste manual por aguja y calibrar cada producto por separado.',
     'Prelavado: '+number(scaled('plant.injector.mChemical',60000/rho,time=18))+'; jabón: '+number(scaled('plant.injector.mChemical',60000/rho,time=33))+'; cera: '+number(scaled('plant.injector.mChemical',60000/rho,time=48))+' L/min.',
     'Dosis estimadas, no garantizadas', 'Instantáneas a 18, 33 y 48 s; no son promedios de etapa. Los químicos se representan como agua: confirmar viscosidad, densidad y concentración real.'),
    ('Tanques y sifones', 'Usar tapas con respiradero y conservar sifones de 12,7 mm interiores.',
     'Extracción desde tanque en jabón: '+number(scaled('plant.soap.tankDraw',60000/rho,time=33))+' L/min; manifold: '+number(pg('plant.chemicalManifoldPressure',33))+' bar(g).',
     'Configuración viable en el modelo', 'Tanques atmosféricos y mangueras cebadas ideales. Longitud de 2 m y elevación de 0,5 m son supuestos; no se comparó esta ejecución con un tanque hermético.'),
    ('Selección y purga', 'Abrir una sola electroválvula y seleccionar baja presión durante la purga previa de 2 s.',
     'Apertura de la NC de jabón: '+number(get('plant.soap.solenoid.opening',33),0)+'; bloqueo de inyección: '+number(get('plant.chemicalInjectionBlocked',33),0)+'.',
     'Lógica comprobada en el ciclo', 'ProgramController y current-validation.json: máximo una rama abierta. La selección ideal representa el cambio de boquilla por el operador; 2 s no certifican limpieza ni eliminación de residuos químicos.'),
    ('Calentador externo', 'Mantener entrada a 13 °C y ajuste manual de gas al máximo; usar temperatura con caudal para evaluar el servicio.',
     'Antes del cierre (55 s): '+number(temp('plant.inletHeater.T',55),2)+' °C. Final sin flujo (64 s): '+number(temp('plant.inletHeater.T',64),2)+' °C.',
     'Configuración térmica simulada', 'plant.inletHeater.T y m_flow. La temperatura a 64 s corresponde a agua retenida. Confirmar potencia útil del HACEB y compatibilidad térmica real de la TRENT.'),
    ('Protección de presión', 'Clasificar componentes de alta presión por el máximo del circuito; no por la presión baja durante aplicación.',
     'Presión retenida con pistola cerrada (58 s): '+number(pg('plant.pumpOutletPressure',58),2)+' bar(g); descarga: '+number(get('plant.outletFlow_L_min',58))+' L/min.',
     'Selección mecánica pendiente', 'plant.pumpOutletPressure. Límite parametrizado de bomba: 130 bar; añadir criterios físicos de transitorios y margen al seleccionar equipos. No se simularon golpes de ariete.'),
    ('Electroválvulas químicas', 'Seleccionar NC aptas para aspiración y verificar aislamiento mediante checks.',
     'Jabón a 33 s: entrada NC '+number(pg('plant.soap.valveInletPressure',33))+'; salida NC '+number(pg('plant.soap.valveOutletPressure',33))+' bar(g).',
     'Especificación por confirmar', 'Presiones normales de simulación, no ratings de compra. Compatibilidad de sellos, diferencial mínimo, vacío admisible y escenarios de fallo de checks deben confirmarse.'),
    ('Validez del diseño', 'Calibrar el prototipo antes de adoptar dosis y presiones como especificaciones finales.',
     'La validación actual comprueba balances, exclusividad, aspiración y cierre.',
     'Validación del modelo, no del equipo', 'verification/results/current-validation.json. Los coeficientes de bomba, inyector, válvulas y Cd son supuestos. El límite de garganta puede actuar durante cambios de boquilla.'),
])
decisions['A2'] = 'Ciclo actual de 0–64 s · Evidencias e implicaciones de diseño · Resultados simulados, pendientes de calibración física'
for col, heading in enumerate(['Equipo / aspecto', 'Decisión de diseño', 'Resultado que la respalda', 'Estado', 'Fuente / limitación'],1):
    decisions.cell(4,col).value = heading
widths = [27, 46, 48, 29, 65]
for col,width in zip('ABCDE',widths):
    decisions.column_dimensions[col].width = width
for row in decisions.iter_rows(min_row=5):
    lines = max(len(textwrap.wrap(str(c.value or ''),max(8,int(widths[i]*.9)))) for i,c in enumerate(row))
    decisions.row_dimensions[row[0].row].height = max(48,lines*16+12)

sha=hashlib.sha256((ROOT/'autolavado.mo').read_bytes()).hexdigest()
sheet('Fuentes','FUENTES DE LOS RESULTADOS',[
    text('Modelo simulado','autolavado.SelfServiceCarWashTest','Fuente: autolavado.mo actual.'),
    text('Identidad del modelo',sha,'SHA-256 de autolavado.mo al generar el informe.'),
    text('Resultados','CurrentCycle_res.csv','verification/current/CurrentCycle_res.csv; simulación nueva de 0 a 64 s.'),
    text('Variables y parámetros','CurrentCycle_init.xml','verification/current/CurrentCycle_init.xml; permite resolver aliases y unidades.'),
    text('Formato de celdas',SOURCE.name,'Estilos de título, encabezados y tabla tomados de la hoja Bomba original.'),
    text('Naturaleza de los valores','Resultados simulados','No son mediciones de un prototipo; sin comparaciones ni fórmulas de Excel.'),
    text('Celdas sin valor','Dato no disponible','No se sustituyen datos ausentes por cero.'),
])
temporary=OUTPUT.with_name('Memoria_Calculo_Hidraulica_Modelica_nuevo.xlsx')
book.save(temporary)
check=openpyxl.load_workbook(temporary)
assert abs(check['Boquilla']['B6'].value-1.15)<1e-12
assert not any(c.data_type=='f' for s in check for row in s for c in row)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==before
try:
    temporary.replace(OUTPUT)
except PermissionError:
    # Excel may hold the earlier report open; retain the validated new copy.
    OUTPUT = temporary
print('Workbook:',OUTPUT)
print('Sheets:',len(check.sheetnames),'End time:',end,'Nozzle mm:',get('plant.nozzleDiameter')*1000)
print('t=9s pressure bar(g):',(get('plant.pumpOutletPressure')-patm)/1e5)
print('Chemical L/min at 18,33,48:',[get('plant.injector.mChemical',t)*60000/rho for t in (18,33,48)])
