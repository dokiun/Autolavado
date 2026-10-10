"""Rebuild the audited workbook from the untouched template and OMC exports.

Run OMC scripts first, then this script, then recalculate with Excel.
No Modelica source is modified. All analytical Excel calculations are formulas.
"""
from pathlib import Path
import csv, json, re, hashlib, datetime, math
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parent.parent
BUILD=ROOT/'verification/build'
ORIGINAL=ROOT/'Memoria_Calculo_Hidraulica_Autolavado.xlsx'
OUT=ROOT/'Memoria_Calculo_Hidraulica_Completa.xlsx'
source=(ROOT/'autolavado.mo').read_text(encoding='utf-8')
sha=hashlib.sha256((ROOT/'autolavado.mo').read_bytes()).hexdigest()
original_sha=hashlib.sha256(ORIGINAL.read_bytes()).hexdigest()
now=datetime.datetime.now(datetime.timezone.utc).isoformat()
wb=openpyxl.load_workbook(ORIGINAL)
changes=[]
def put(sheet,cell,value,reason):
    old=wb[sheet][cell].value
    wb[sheet][cell]=value
    wb[sheet][cell].comment=Comment(f'Auditoría: {reason}\nOriginal: {old}', 'Auditoría hidráulica')
    changes.append((sheet,cell,old,value,reason))
def sheet(name,headers):
    s=wb.create_sheet(name);s.append(headers);s.freeze_panes='B2'
    return s
def line_of(text):
    i=source.find(text)
    return source[:i].count('\n')+1 if i>=0 else None
def cite(cls):return f'autolavado.mo:{line_of("model "+cls) or line_of("block "+cls) or line_of("partial model "+cls)} / {cls}'

# Inspecting every original sheet precedes all edits. Keep all supplied constants.
inspection={s.title:{'rows':s.max_row,'cols':s.max_column,'formulas':sum(c.data_type=='f' for row in s for c in row),'values':sum(c.value is not None for row in s for c in row)} for s in wb}
for s in wb:
    for row in s:
        for c in row:
            if c.data_type=='f' and 'Manguera_Pistola!' in c.value and 'B$8' in c.value:
                put(s.title,c.coordinate,c.value.replace('Manguera_Pistola!$B$8','Manguera_Pistola!$B$9'),'La pérdida está en B9; B8 es referencia comercial/texto.')
put('Datos_base','B9',13,'Temperatura de entrada confirmada por el usuario; original 15 °C conservado en registro.')
put('Datos_base','D9','Condición de frontera suministrada','13 °C suministrados, no medidos por este trabajo.')
put('Datos_base','B23','=IF(B12="","",B12)','No suponer calentamiento nulo cuando la salida del calentador es desconocida.')
put('Datos_base','E23','PENDIENTE sin temperatura de salida; el calentador está instalado. Ver Calentador_Modelo para la estimación del modelo.','Separar diseño real del modelo aproximado.')
put('Bomba','D6','Dato comercial preliminar NO verificado','No se localizó manual/placa oficial de referencia exacta.')
put('Bomba','D7','Dato comercial preliminar NO verificado','Sin curva que confirme presión y caudal simultáneos.')
put('Bomba','D8','Dato comercial preliminar NO verificado','Discrepa con referencia HLT395-1800 del modelo.')
# References to blank cells must not manufacture zero pressure/NPSH.
for s,addr,target in [('Resumen','B23','Alimentacion!B12'),('Resumen','B24','Alimentacion!B14'),('Estaciones','E7','Alimentacion!B12'),('Estaciones','F7','Datos_base!B23')]:
    put(s,addr,f'=IF({target}="","",{target})','Propagar pendientes sin convertir blanco en cero.')
put('Propiedades','B6','=IF(Datos_base!B23="","",Datos_base!B23)','Sin temperatura de succión identificada no evaluar propiedades a cero grados.')
for addr in ['F8','F9','F10','F11','F12','F13']:
    put('Estaciones',addr,'=IF(Datos_base!B23="","",Datos_base!B23)','Propagar temperatura pendiente sin asumir cero grados.')
put('Boquilla','B12','=IF(OR(B6="",B7="",B8="",B8<=0,B10="",B10<=0),"",((B7/2)*(B6/60000/(B8*B10))^2)/100000)','Control Cd positivo y denominadores.')
put('Boquilla','B7','=IF(Propiedades!B12="","",Propiedades!B12)','Propiedad de mezcla pendiente si temperatura succión no definida.')
put('Bomba','B17','=IF(OR(B16="",B6="",B6<=0),"",B16/B6)','Evitar división por cero al editar potencia.')
for r in range(8,28):
    put('Tuberias',f'R{r}',f'=IF(OR(P{r}="",P{r}<=0,H{r}="",H{r}<=0,I{r}="",I{r}<0),"",IF(P{r}<2300,64/P{r},IF(P{r}<4000,"",0.25/(LOG10(I{r}/(3.7*H{r})+5.74/(P{r}^0.9))^2))))','Transición 2300–4000 sin correlación validada: dejar resultado pendiente; Darcy, no Fanning.')
put('LEEME','B14','Propiedades originales conservadas como aproximaciones. Modelo MSL: rho, mu y cp constantes. Presión de vapor auditada con IAPWS en NPSH_Modelo.','Distinguir propiedades de plantilla y propiedades del modelo.')
for r in range(8,28):
    # Completeness includes K, properties, and transition result, not just geometry.
    put('Tuberias',f'V{r}',f'=IF(B{r}="","",IF(OR(G{r}="",H{r}="",H{r}<=0,I{r}="",J{r}="",K{r}="",L{r}="",M{r}="",M{r}<=0,AND(K{r}<>0,R{r}="")),"INCOMPLETO","COMPLETO"))','No aceptar K desconocido ni régimen transicional como pérdida cero.')
put('Tuberias','B36','=IF(B35=0,"GEOMETRÍA DECLARADA PARA TRAMOS ACTIVOS","PENDIENTE: completar geometría, K, propiedades y régimen")','Completitud física, además de geométrica.')
for r in range(31,35):
    typ=['Succión','Antes Venturi','Después Venturi','Químicos'][r-31]
    put('Tuberias',f'B{r}',f'=IF(COUNTIFS($B$8:$B$27,"{typ}",$V$8:$V$27,"INCOMPLETO")>0,"",SUMIF($B$8:$B$27,"{typ}",$U$8:$U$27))','No sumar incógnitas como cero.')

meta=sheet('Auditoria',['Concepto','Valor','Estado / interpretación'])
for row in [('Fecha UTC',now,'Ejecución del generador'),('SHA256 Modelica',sha,'Fuente sin modificaciones'),('SHA256 original',original_sha,'Original conservado'),('Modelo','autolavado.SelfServiceCarWashTest','631 ecuaciones / 631 variables; checkModel'),('MSL','4.1.0','ConstantPropertyLiquidWater'),('OMC','1.27.1','DASSL; tol 1e-6'),('Diseño real','PENDIENTE','No validado experimentalmente'),('Red Bogotá y filtro','No representados en Modelica','La frontera es 101325 Pa absolutos'),('Referencia bomba Excel','2100 W / 110 bar nominal / 150 bar máximo','Comercial preliminar; no curva'),('Referencia bomba modelo','1800 W / 130 bar máximo / carga 100 bar','Datos previos usuario + supuestos'),('Manguera','Excel 4 m; modelo 8 m','No se sustituyen entre sí'),('Régimen','Instantáneas y transitorios','No llamar estacionarios a todos los puntos'),('Guía','Analitica_Modelo: Q, contrapresión y estado editables','Presupuesto condicional, no solucionador completo de red'),('Resultado experimental','Ninguno disponible','Consistencia de simulación no valida prototipo')]:meta.append(row)
for name,stats in inspection.items():meta.append((f'Inspección {name}',json.dumps(stats),'Todas las hojas originales conservadas'))

params=sheet('Parametros_Modelica',['Ruta / variable','Valor resuelto / expresión','Unidad','Tipo de dato','Fuente','Estado','Descripción'])
instance=json.loads((ROOT/'verification/model-instance.json').read_text(encoding='utf-8'))
components=[];connections=[];pmap={}
def unit_of(t):
    if not isinstance(t,dict):return ''
    for e in t.get('elements',[]):
        if e.get('$kind')=='extends':
            u=e.get('modifiers',{}).get('unit','')
            if isinstance(u,dict):u=u.get('$value','')
            if u:return str(u).strip('"')
            u=unit_of(e.get('baseClass'))
            if u:return u
    return ''
def walk(obj,path=''):
    for e in obj.get('elements',[]):
        if e.get('$kind')=='extends' and isinstance(e.get('baseClass'),dict):
            walk(e['baseClass'],path)
        if e.get('$kind')!='component':continue
        p=f'{path}.{e["name"]}' if path else e['name'];t=e.get('type');typ=t.get('name') if isinstance(t,dict) else t
        pref=e.get('prefixes',{});val=e.get('value',{});mods=e.get('modifiers',{})
        if pref.get('variability') in ('parameter','constant'):
            v=val.get('value',val.get('binding','PENDIENTE'))
            if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
            unit=(mods.get('unit','') if isinstance(mods,dict) else '') or unit_of(t)
            comment=e.get('comment','')
            kind='Parámetro supuesto'
            if any(x in p for x in ['diameter','Diameter','length','Length','roughness','Volume','crossSection','height']):kind='Parámetro geométrico (supuesto salvo nota)'
            if any(x in p for x in ['rho','cp_const','eta_const','d_const']):kind='Propiedad física'
            if p in ['plant.pAtmosphere','plant.supplyTemperature']:kind='Condición de frontera'
            params.append((p,v,str(unit),kind,str(obj.get('source',{})), 'Modelado; no validado experimentalmente',comment));pmap[p]=params.max_row
        if isinstance(t,dict):
            if t.get('restriction') in ('model','block'):components.append((p,typ,t.get('comment',''),str(t.get('source',{}))))
            if t.get('restriction') not in ('connector','type','package','function'):walk(t,p)
    for c in obj.get('connections',[]):connections.append((path,json.dumps(c,ensure_ascii=False)))
walk(instance)
hier=sheet('Componentes',['Ruta','Clase','Función','Fuente instanciada','Estado','Entradas','Salidas','Ecuaciones / hipótesis'])
for p,t,comment,src in components:
    state='Aproximado' if t.startswith('autolavado.') else 'MSL; parametrización pendiente de ensayo'
    hier.append((p,t,comment,src,state,'Ver Parametros_Modelica y Conexiones','Ver Resultados_OpenModelica','Ver Ecuaciones; medio agua equivalente'))
hier.append(('plant.filtro','AUSENTE','Filtro conceptual no implementado','autolavado.SelfServiceCarWash','Pendiente','Curva Δp–Q / obstrucción','Pérdida de presión','Sin ecuación constitutiva en fuente'))
con=sheet('Conexiones',['Ámbito','Conexión instanciada'])
for row in connections:con.append(row)
decl=sheet('Declaraciones',['Clase','Tipo','Nombre','Unidad declarada / tipo','Expresión / declaración','Línea','Estado'])
cls='autolavado'
for ln,txt in enumerate(source.splitlines(),1):
    m=re.match(r'\s*(?:partial )?(model|block|function) (\w+)',txt)
    if m:cls=m[2]
    m=re.match(r'\s*((?:final )?parameter |constant |output )?((?:SI\.|Modelica\.)?[\w.]+) (\w+(?:\[[^;]*?\])?)(.*);',txt)
    if m and not txt.lstrip().startswith(('connect','end','annotation','import','extends','assert')):
        decl.append((cls,(m[1] or 'Variable calculada').strip(),m[3],m[2],txt.strip(),ln,'No implica validación física'))

# Compact, editable SI parameters for analytical model, independent of manufacturer claims.
inp=sheet('Entradas_Modelo',['ID','Valor','Unidad','Tipo de dato','Fuente / ruta','Estado'])
defs=[('rho',995.586,'kg/m3','Propiedad física','MSL Media/Water/package.mo:52','Constante del modelo'),('mu',.001,'Pa.s','Propiedad física','MSL Media/Water/package.mo:53','Constante del modelo'),('cp',4184,'J/(kg.K)','Propiedad física','MSL Media/Water/package.mo:50','Constante del modelo'),('g',9.80665,'m/s2','Propiedad física','Modelica.Constants.g_n','Estándar'),('patm',101325,'Pa(abs)','Condición de frontera','plant.pAtmosphere','Modelo; no Bogotá medida'),('Tin',13,'°C','Condición de frontera','plant.supplyTemperature - 273.15','Usuario'),('Qref',7/60000,'m3/s','Parámetro fabricante preliminar','plant.referenceFlow','Usuario; sin curva'),('Lhose',8,'m','Parámetro geométrico','plant.hoseLength','Modelo; diseño Excel 4 m'),('Dhose',.008,'m','Parámetro supuesto','plant.hoseDiameter','Por medir'),('epshose',2.5e-5,'m','Parámetro supuesto','plant.hose.roughness','Por medir'),('dploss',180000,'Pa','Parámetro supuesto','plant.injector.dpLossNominal','Calibrar AQB'),('recovery',180000,'Pa','Parámetro supuesto','plant.injector.recoveryNominal','Calibrar AQB'),('pfloor',20000,'Pa(abs)','Parámetro supuesto','plant.injector.pFloor','No modela cavitación'),('qchemref',.25/60000,'m3/s','Parámetro supuesto','plant.injector.mChemicalNominal/rho','Calibrar'),('dpchem',5000,'Pa','Parámetro supuesto','plant.injector.dpCheckNominal','Calibrar'),('pcrack',1000,'Pa','Parámetro supuesto','plant.injector.pCrack','Calibrar'),('mreg',1e-6,'kg/s','Parámetro supuesto','plant.injector.mRegularization','Regularización numérica'),('Pref',100e5,'Pa','Parámetro supuesto','plant.workingPressure','Carga de referencia'),('slip',.08,'1','Parámetro supuesto','plant.bomba.slipFraction','Sin curva'),('Pel',1800,'W','Parámetro fabricante preliminar','plant.bomba.ratedElectricalPower','Usuario; placa pendiente'),('eta',.7,'1','Parámetro supuesto','plant.bomba.hydraulicEfficiency','Global eléctrica→hidráulica'),('Pgas',12000,'W','Parámetro fabricante preliminar','plant.inletHeater.ratedPower','Gas entrada supuesto'),('etaheat',.85,'1','Parámetro supuesto','plant.inletHeater.efficiency','Calibrar'),('gasSetting',1,'1','Condición de frontera','plant.inletHeater.gasSetting','Manual máximo'),('UA',800,'W/K','Parámetro supuesto','plant.inletHeater.waterConductance','Calibrar'),('Uamb',2,'W/K','Parámetro supuesto','plant.inletHeater.ambientConductance','Calibrar'),('Tamb',20,'°C','Parámetro supuesto','plant.inletHeater.ambientTemperature','Modelo'),('Qheatref',8/60000,'m3/s','Parámetro fabricante preliminar','plant.inletHeater.ratedFlow','No impone caudal'),('dpheat',20000,'Pa','Parámetro supuesto','plant.inletHeater.nominalPressureDrop','Calibrar'),('dpGun',20000,'Pa','Parámetro supuesto','plant.gun.pistola.dp_nominal','Pistola abierta a Qref'),('dpLow',40000,'Pa','Parámetro supuesto','plant.gun.dpLow','Boquilla LP'),('dpHigh',97e5,'Pa','Parámetro supuesto','plant.gun.dpHigh','Boquilla HP'),('Lchem',2,'m','Parámetro supuesto','plant.chemicalSiphonLength','Por medir'),('Dchem',.0127,'m','Parámetro geométrico','plant.chemicalSiphonDiameter','Interior confirmado usuario'),('epschem',2.5e-6,'m','Parámetro supuesto','plant.soap.siphon.roughness','Por medir'),('lift',.5,'m','Parámetro supuesto','plant.chemicalSuctionLift','Por medir'),('Vtank',.01,'m3','Parámetro supuesto','plant.soap.tankLiquidVolume','Por medir'),('Atank',.03,'m2','Parámetro supuesto','plant.soap.tankCrossSection','Por medir'),('qbranch',.20/60000,'m3/s','Parámetro supuesto','plant.soap.mNominal/rho','Cera usa .10 L/min'),('dpNC',4000,'Pa','Parámetro supuesto','plant.soap.solenoid.dpNominal','12 V NC acción directa equivalente'),('dpCheck',1000,'Pa','Parámetro supuesto','plant.soap.checkValve.dpNominal','Calibrar'),('dpNeedle',42000,'Pa','Parámetro supuesto','plant.soap.needle.dpNominal','Calibrar'),('needle',1,'1','Condición de frontera','plant.soap.needleSetting','Ajuste manual'),('Dsuccion',None,'m','Parámetro geométrico','Brida de succión real','PENDIENTE'),('NPSHr',None,'m','Parámetro fabricante','Fabricante TRENT exacto','PENDIENTE'),('etaPump',None,'1','Parámetro fabricante','Eficiencia hidráulica eje→agua','PENDIENTE'),('rpm',None,'rpm','Parámetro fabricante','Velocidad del eje real','PENDIENTE'),('Cd',None,'1','Parámetro fabricante','Boquilla exacta','PENDIENTE'),('pDesign',None,'Pa(g)','Parámetro fabricante','Presión admisible mínima conjunto','PENDIENTE')]
defs += [('Vwater',.0003,'m3','Parámetro supuesto','plant.inletHeater.waterVolume','Calibrar'),('Cm',1500,'J/K','Parámetro supuesto','plant.inletHeater.exchangerHeatCapacity','Calibrar'),('drawnVolume',0,'m3','Condición de frontera','Ensayo condicional: tanque inicialmente lleno','No resultado de simulación'),('uNC',1,'1','Condición de frontera','Ensayo condicional rama abierta','Editable; sólo rama seleccionada')]
R={}
for d in defs:inp.append(d);R[d[0]]=inp.max_row
def P(k):return f"Entradas_Modelo!$B${R[k]}"

eqs=[('E01','Continuidad','m=ρQ; Q=A v; Σm=dm/dt','TwoPort','Medio de densidad constante'),('E02','Tubería','A=πD²/4; Re=ρ|v|D/μ; Δp=f(L/D)ρv²/2+ΣKρv²/2+ρgΔz','ChemicalBranch','MSL StaticPipe, gravedad; sin almacenamiento'),('E03','Fricción','f=64/Re laminar; f=.25/[log10(ε/(3.7D)+5.74/Re^.9)]² turbulento','SelfServiceCarWash','MSL Detailed: from_dp=true usa inversión Colebrook; fórmula directa Swamee–Jain, transición cúbica'),('E04','Calentador hidráulico','Δp=Δpnom (Q/Qnom) sqrt((Q/Qnom)²+1e-8)','HacebGasHeater','Resistencia supuesta, no filtro'),('E05','Calentador energía','Cm dTm/dt=ηPgas-UA(Tm-T)-Uamb(Tm-Tamb); ρVcp dT/dt=Σm(hstream-h)+UA(Tm-T)','HacebGasHeater','Dos estados térmicos, agua y metal; sensores externos'),('E06','Bomba','Q=min(Qmax,Qref max(0,(1+s)ω-sΔp/Pref),ηPelω/max(1,Δp))','PressureWasherPump','ω normalizada, sin rpm ni cavitación; motor primer orden'),('E07','Potencia','Ph=QΔp; H=Δp/(ρg); Peje=Ph/ηp','PressureWasherPump','ηp real desconocida; η=.7 del modelo es global'),('E08','Alivio','m=mnom positiveRoot((Δp-pset)/banda)','UnloaderValve','pset=117 bar(g); sin histéresis ni unloader real'),('E09','Venturi principal','Δp=Δpnom r sqrt(r²+.001²); pt=max(pFloor,pout-recovery r²)','DownstreamInjector','Sin geometría ni curva AQB'),('E10','Aspiración','mc=mcNom positiveRoot((pm-pt-pcrack)/dpCheck) max(0,mw)/sqrt(mw²+mreg²)','DownstreamInjector','Antirretorno HP ideal; pFloor no es vapor'),('E11','Mezcla','mout=mw+mc; hout=(mw hw+mc hc+mreg hw)/(mw+mc+mreg)','DownstreamInjector','Regularización y agua equivalente; error entalpía pequeño no exactamente cero'),('E12','Tanque','V=V0-Vextraído; pGas=patm; pSalida=patm+ρgV/Atank','ChemicalTank','Tapa con respiradero; sin aire en manguera ni cebado'),('E13','NC','m=mnom u regRoot(Δp/dpnom,.001)','ChemicalSolenoid','u=0 cierre ideal sin fuga; regRoot(x,δ)=x/(x²+δ²)^.25'),('E14','Check rama','m=mnom positiveRoot((Δp-pcrack)/dpnom)','ChemicalCheckValve','Bloqueo inverso ideal salvo ensayo failedOpen'),('E15','Aguja','m=mnom setting regRoot(Δp/dpnom,.001)','ChemicalNeedleValve','Bidireccional, sin geometría real'),('E16','Elasticidad','V=V0+C(p-p0); ρC dp/dt=Σm; ρVcp dT/dt=Σm(hstream-h)','CompliantVolume','Omite trabajo de compresión; no golpe de ariete'),('E17','Pistola / boquilla','m=opening Av sqrt(ρ) regRoot2(Δp,dp_small,...); CdA=Av/√2','SprayGun','MSL ValveIncompressible, característica lineal, antirretorno'),('E18','Concentración','Mman dxi/dt=mProducto-mSuma xi; ρVtracer dyi/dt=mc xi-mout yi','SelfServiceCarWash','Trazador pasivo, no composición Medium; no mezcla química real'),('E19','NPSHA','(ps,abs-pv)/(ρg)+vs²/(2g)','SelfServiceCarWash','Succión brida; sin NPSHr no hay criterio validado'),('E20','Presión de vapor','pv=Pc exp[(Tc/T)Σai(1-T/Tc)^bi]','SelfServiceCarWash','IAPWS SR1-86(1992), auxiliar Excel; NO implementado en OMC'),('E21','Control','exclusión químicos; purga 2s solo con Q>2 L/min; LP primer orden .2s','ProgramController','Autostop por gatillo, no sensor real de presión'),('E22','Raíz unilateral','0 si x≤0; √δ[2.5(x/δ)²-1.5(x/δ)³] si 0<x<δ; √x resto','DownstreamInjector','positiveRoot δ=.001; continuidad C1')]
eq=sheet('Ecuaciones',['ID','Subsistema','Ecuación','Variables / unidades','Hipótesis','Fuente OpenModelica','Referencia bibliográfica'])
for eid,sub,formula,cl,hyp in eqs:eq.append((eid,sub,formula,'SI; p absoluta en puertos; dp diferencia',hyp,cite(cl),'MSL 4.1.0 fuente local; IAPWS si E20'))
eq.append(('E23','Filtro','PENDIENTE Δp(Q,carga)','Pa, m3/s','Ausente de Modelica','No existe','Fabricante / ensayo requerido'))

# Read valid exports with aliases. Every exported value comes from OMC, never from guesses.
def load_run(name):
    path=BUILD/f'{name}_res.csv'
    if not path.exists():raise RuntimeError(f'Resultado necesario ausente: {path}')
    xml={v.get('name'):v for v in ET.parse(BUILD/f'{name}_init.xml').findall('.//ScalarVariable')}
    class Row(dict):
        def __missing__(self,k):
            v=xml[k]
            if v.get('alias') in ('alias','negatedAlias'):
                a=self[v.get('aliasVariable')];return -a if v.get('alias')=='negatedAlias' else a
            t=next(iter(v));s=t.get('start')
            if s is not None and v.get('variability') in ('parameter','constant'):return float({'true':1,'false':0}.get(s,s))
            raise KeyError(k)
    with path.open(encoding='utf-8-sig') as f:rows=[Row({k:float(v) for k,v in r.items()}) for r in csv.DictReader(f)]
    return rows,xml,datetime.datetime.fromtimestamp(path.stat().st_mtime,datetime.timezone.utc).isoformat()
runs={n:load_run(n) for n in ['InjectionCycle','InjectionVented','InjectionProtected','InjectionFailed','HacebPlant','HacebFull','HacebFlame','HacebVent','HacebOverheat','HacebLowFlow','HacebPower']}
for k,v in runs['InjectionCycle'][1].items():
    if v.get('variability') not in ('parameter','constant'):continue
    scalar=next(iter(v));unit=scalar.get('unit','1')
    if k not in pmap:
        params.append((k,'PENDIENTE' if scalar.get('start') is None else scalar.get('start'),unit,'Parámetro del modelo','','',''));pmap[k]=params.max_row
    rr=pmap[k];params.cell(rr,3,unit)
    params.cell(rr,5,f'{v.get("fileName")}:{v.get("startLine")} / {k}')
    if v.get('causality')=='calculatedParameter':
        params.cell(rr,4,'Parámetro calculado')
        params.cell(rr,6,'Binding / ecuación; no es dato de fabricante')
    if params.cell(rr,2).value=='PENDIENTE' and v.get('causality')=='calculatedParameter':params.cell(rr,6,'Calculado por inicialización OMC; no exportado en CSV. Ver ecuación fuente.')
required_ends={'InjectionCycle':64,'InjectionVented':120,'InjectionProtected':15,'InjectionFailed':15,'HacebPlant':30,'HacebFull':120,'HacebFlame':40,'HacebVent':40,'HacebOverheat':120,'HacebLowFlow':40,'HacebPower':120}
for n,(rows,_,_) in runs.items():
    assert abs(rows[-1]['time']-required_ends[n])<1e-6,(n,'incomplete')
    assert all(math.isfinite(v) for r in rows for v in r.values()),(n,'nonfinite')
raw=sheet('Resultados_OpenModelica',['Escenario','Fecha UTC','t [s]','Modo','Variable','Valor','Unidad','Tipo de dato','Estado / fuente'])
mapping=sheet('Mapa_Modelica_Excel',['Variable física','Símbolo','Variable OpenModelica','Componente','Unidad','Celda Excel','Ecuación','Tipo de dato','Fuente','Validación'])
export=[];rawrefs={};snapshots=[];csvunits={}
selected=['pumpInletPressure','pumpOutletPressure','injectorInletPressure','injectorOutletPressure','throatPressure','chemicalManifoldPressure','gunInletPressure','hosePressureDrop','injectorPressureDrop','nozzlePressureDrop','waterFlow_L_min','outletFlow_L_min','chemicalFlow_L_min','bypassFlow_L_min','pumpFlow_L_min','waterTemperature','outletTemperature','heatingPower','mainMassResidual','chemicalMassResidual','injectorMassResidual','chemicalFractionAtInjector','outletChemicalFraction','inletHeater.T','inletHeater.metalTemperature','inletHeater.gasPower','inletHeater.heatToWater','inletHeater.ambientLoss','inletHeater.energyResidual','bomba.hydraulicPower','bomba.speed','gun.pistola.dp','gun.nozzle.opening','gun.nozzle.Av','gun.pistola.Av','gun.nozzle.rho_nominal','injector.mWater','injector.mChemical','injector.mMixture','injector.mixture.h_outflow','injector.water.h_outflow','injector.chemical.h_outflow','accumulator.storageRate','chemicalManifold.storageRate','der(inletHeater.T)','der(inletHeater.metalTemperature)']
for b in ['prewash','soap','wax']:
    selected += [f'{b}.{v}' for v in ['tankDraw','m_flow','tank.liquidVolume','tank.gasPressure','valveInletPressure','valveOutletPressure','valveDifferentialPressure','solenoid.u','solenoid.m_flow','checkValve.m_flow','checkValve.dp','needle.m_flow','needle.dp','siphon.port_a.p','siphon.port_b.p','beforeCheck.storageRate','afterCheck.storageRate']]
selected += ['accumulator.h','chemicalManifold.h','chemicalManifold.mass','hose.flowModel.vs[1]']
for j in [1,2,3]:
    selected += [f'injectedProductFlow[{j}]',f'manifoldConcentration[{j}]',f'outletConcentration[{j}]',f'der(outletConcentration[{j}])',f'der(manifoldConcentration[{j}])']
def fullvar(v):return 'der(plant.'+v[4:] if v.startswith('der(') else 'plant.'+v
for name,(rows,xml,date) in runs.items():
    plant=name in ('InjectionCycle','InjectionVented','HacebPlant')
    keys=[fullvar(v) for v in selected] if plant else ([k for k in xml if k.startswith('branch.') and any(k.endswith(s) for s in ['tankDraw','valveInletPressure','valveOutletPressure','valveDifferentialPressure','solenoid.m_flow'])] if name.startswith('Injection') else ['heater.'+v for v in ['T','metalTemperature','m_flow','gasPower','heaterPower','heatToWater','ambientLoss','energyResidual','burnerLit','safetyLocked']])
    keys=list(dict.fromkeys(k for k in keys if k in xml or k in rows[0]))
    # CSV contains full time histories; workbook holds snapshots for manageable size.
    for r in rows:
        mode=str(int(r.get('plant.selectedProgram',-1))) if plant else 'banco'
        csvrow={'escenario':name,'fecha_UTC':date,'tiempo_s':r['time'],'modo_programa':mode,'tipo_dato':'Resultado simulado','sha256_modelo':sha}
        for k in keys:
            try:v=r[k]
            except KeyError:continue
            unit=next(iter(xml[k])).get('unit','1') if k in xml else '1'
            csvrow[k]=v;csvunits[k]=unit
        export.append(csvrow)
    times=([0,2,4,9,10,18,20,25,33,35,40,48,50,55,56,57,58,60,64] if name=='InjectionCycle' else [rows[-1]['time']] if name in ['InjectionVented','HacebPlant','HacebFull','HacebLowFlow','HacebPower'] else [0,2,5,9,14,19,26,rows[-1]['time']])
    for t in times:
        r=min(rows,key=lambda r:abs(r['time']-t));mode=str(int(r.get('plant.selectedProgram',-1))) if plant else 'banco'
        if plant:snapshots.append((name,r['time'],mode,r))
        for k in keys:
            try:v=r[k]
            except KeyError:continue
            unit=next(iter(xml[k])).get('unit','1') if k in xml else '1'
            raw.append((name,date,r['time'],mode,k,v,unit,'Resultado simulado','Instantánea; no estacionario garantizado'))
            rawrefs[(name,r['time'],k)]=f"Resultados_OpenModelica!F{raw.max_row}"
            if t==times[-1]:mapping.append((k,k,k,'.'.join(k.split('.')[:-1]),unit,f'Resultados_OpenModelica!F{raw.max_row}','Ver Ecuaciones','Resultado simulado',f'{name}_res.csv; SHA fuente {sha}','Simulado; no experimental'))
with (ROOT/'Resultados_Simulacion.csv').open('w',encoding='utf-8-sig',newline='') as f:
    cols=['escenario','fecha_UTC','tiempo_s','modo_programa','tipo_dato','sha256_modelo']+list(csvunits)
    wr=csv.writer(f);wr.writerow([c if c not in csvunits else f'{c} [{csvunits[c]}]' for c in cols]);wr.writerows([[row.get(c,'') for c in cols] for row in export])
for row in defs:mapping.append((row[0],row[0],row[4],'plant',row[2],f'Entradas_Modelo!B{R[row[0]]}','Ver Ecuaciones',row[3],row[4],row[5]))
for label,var,unit,cell,typ,status in [
    ('Temperatura fuente','plant.supplyTemperature','°C','Datos_base!B9','Condición de frontera','13 °C suministrados'),
    ('Atmósfera local','plant.pAtmosphere','kPa(abs)','Datos_base!B7','Parámetro supuesto','74 kPa diseño;101.325 kPa modelo; medir'),
    ('Caudal comercial','plant.referenceFlow','L/min','Datos_base!B10','Parámetro fabricante preliminar','7 L/min; sin curva'),
    ('Potencia eléctrica','plant.bomba.ratedElectricalPower','W','Bomba!B6','Parámetro fabricante preliminar','2100 Excel;1800 modelo'),
    ('Presión nominal','plant.workingPressure','bar(g)','Bomba!B7','Parámetro fabricante preliminar','110 Excel;100 supuesto modelo'),
    ('Presión máxima','plant.maximumPressure','bar(g)','Bomba!B8','Parámetro fabricante preliminar','150 Excel;130 modelo'),
    ('NPSHr','AUSENTE','m','Bomba!B14','Parámetro fabricante','PENDIENTE'),
    ('Pérdida filtro','AUSENTE','kPa','Alimentacion!B6','Parámetro fabricante','PENDIENTE; no componente OMC'),
    ('Pérdida calentador','plant.inletHeater.port_a.p-plant.inletHeater.port_b.p','kPa','Alimentacion!B7','Parámetro fabricante','Real pendiente; modelo20 kPa a8 L/min supuesto'),
    ('Succión brida','plant.pumpInletPressure','kPa(abs)','Alimentacion!B12','Variable calculada','Frontera distinta; no copiar resultado modelo'),
    ('NPSHA','Auxiliar; AUSENTE OMC','m','Alimentacion!B14','Variable calculada','PENDIENTE diámetro/p/NPSHr criterio'),
    ('Contrapresión Venturi','plant.injectorOutletPressure','bar(g)','Venturi!B9','Parámetro fabricante','Curva real PENDIENTE'),
    ('Pérdida Venturi','plant.injectorPressureDrop','kPa','Venturi!B10','Parámetro fabricante','Real pendiente;coeficiente modelo supuesto'),
    ('Caudal químico','plant.injector.mChemical/rho','L/min','Venturi!B11','Parámetro fabricante / experimental requerido','No sobrescribir con simulación'),
    ('Densidad producto','Medium.d_const equivalente','kg/m3','Venturi!B16','Propiedad física','Real PENDIENTE'),
    ('Viscosidad producto','Medium.eta_const equivalente','Pa.s','Venturi!B17','Propiedad física','Real PENDIENTE'),
    ('Garganta','plant.throatPressure','bar(g)','Venturi!B18','Parámetro fabricante / experimental requerido','Geometría/curva PENDIENTE'),
    ('Manguera longitud','plant.hoseLength','m','Tuberias!G12','Parámetro geométrico','4 m diseño;8 m modelo'),
    ('Manguera diámetro','plant.hoseDiameter','mm','Tuberias!H12','Parámetro geométrico','Real PENDIENTE;8 mm supuesto'),
    ('Sifón diámetro','plant.chemicalSiphonDiameter','mm','Tuberias!H14','Parámetro geométrico','12.7 mm interior confirmado; no geometría global definida'),
    ('Pérdida pistola','plant.gun.pistola.dp','kPa','Manguera_Pistola!B9','Parámetro fabricante / experimental requerido','Real PENDIENTE'),
    ('Coeficiente descarga','No separado de Av','1','Boquilla!B8','Parámetro supuesto','0.8 plantilla;no identificado por OMC'),
    ('Diámetro boquilla','No separado de Av','mm','Boquilla!B9','Parámetro geométrico','PENDIENTE'),
]:mapping.append((label,label,var,'plant',unit,cell,'Ver Ecuaciones',typ,'Plantilla + fuente Modelica auditada',status))

# Constitutive workbook: independent inputs Q, backpressure, manifold, gun/nozzle position.
a=sheet('Analitica_Modelo',['Variable','Valor','Unidad','Ecuación / interpretación','Tipo de dato'])
ar={}
def A(k):return f"Analitica_Modelo!$B${ar[k]}"
def add(k,v,u,n,typ='Variable calculada'):
    a.append((k,v,u,n,typ));ar[k]=a.max_row
add('Q',f'={P("Qref")}','m3/s','Caudal de ensayo editable; no punto de operación certificado','Condición de frontera')
add('Qquim',None,'m3/s','PENDIENTE: dato de fabricante o ensayo; ensayo condicionado editable','Condición de frontera')
add('pSalidaVenturi',None,'Pa(abs)','Contrapresión editable; necesaria para cerrar garganta','Condición de frontera')
add('pManifold',None,'Pa(abs)','Presión colector; depende de ramas y almacenamiento','Condición de frontera')
add('omega',1,'1','Motor normalizado, NO rpm','Condición de frontera')
add('uGun',1,'1','Pistola abierta 1 / cerrada 0','Condición de frontera')
add('uLP',0,'1','Boquilla HP=0, LP=1','Condición de frontera')
add('Qmezcla',f'=IF({A("Qquim")}="","",{A("Q")}+{A("Qquim")})','m3/s','E11')
add('Area',f'=PI()*{P("Dhose")}^2/4','m2','E02')
add('v',f'=IF({A("Qmezcla")}="","",{A("Qmezcla")}/{A("Area")})','m/s','E02')
add('Re',f'=IF({A("v")}="","",{P("rho")}*ABS({A("v")})*{P("Dhose")}/{P("mu")})','1','E02')
add('f',f'=IF(OR({A("Re")}="",{A("Re")}<=0),"",IF({A("Re")}<2300,64/{A("Re")},IF({A("Re")}<4000,"",.25/(LOG10({P("epshose")}/(3.7*{P("Dhose")})+5.74/{A("Re")}^.9)^2))))','1','E03 Swamee–Jain; transición pendiente')
add('dpHose',f'=IF({A("f")}="","",{A("f")}*{P("Lhose")}/{P("Dhose")}*{P("rho")}*{A("v")}^2/2)','Pa','E02; compara con MSL inversa Colebrook')
add('dpVenturi',f'={P("dploss")}*({A("Q")}/{P("Qref")})*SQRT(({A("Q")}/{P("Qref")})^2+.001^2)','Pa','E09')
add('pEntradaVenturi',f'=IF({A("pSalidaVenturi")}="","",{A("pSalidaVenturi")}+{A("dpVenturi")})','Pa(abs)','E09')
add('pGarganta',f'=IF({A("pSalidaVenturi")}="","",MAX({P("pfloor")},{A("pSalidaVenturi")}-{P("recovery")}*({A("Q")}/{P("Qref")})^2))','Pa(abs)','E09; pFloor NO validación cavitación')
add('xAspiracion',f'=IF(OR({A("pManifold")}="",{A("pGarganta")}=""),"",({A("pManifold")}-{A("pGarganta")}-{P("pcrack")})/{P("dpchem")})','1','E10')
add('qAspirado',f'=IF({A("xAspiracion")}="","",{P("qchemref")}*IF({A("xAspiracion")}<=0,0,IF({A("xAspiracion")}<.001,SQRT(.001)*(2.5*({A("xAspiracion")}/.001)^2-1.5*({A("xAspiracion")}/.001)^3),SQRT({A("xAspiracion")})))*MAX(0,{P("rho")}*{A("Q")})/SQRT(({P("rho")}*{A("Q")})^2+{P("mreg")}^2))','m3/s','E10; predicción condicional, rama debe cerrar balance')
add('fraccion',f'=IF(OR({A("Qmezcla")}="",{A("Qmezcla")}<=0),"",{A("Qquim")}/{A("Qmezcla")})','1','Fracción volumétrica=masica SOLO igual densidad')
add('openingNozzle',f'=SQRT({P("dpLow")}/{P("dpHigh")})+(1-SQRT({P("dpLow")}/{P("dpHigh")}))*{A("uLP")}','1','E17')
add('dpGun',f'=IF(OR({A("Qmezcla")}="",{A("uGun")}<=0),"",{P("dpGun")}*({A("Qmezcla")}/{P("Qref")})^2/{A("uGun")}^2)','Pa','Aproximación turbulenta; cierre requiere transitorio')
add('dpNozzle',f'=IF({A("Qmezcla")}="","",{P("dpLow")}*({A("Qmezcla")}/{P("Qref")})^2/{A("openingNozzle")}^2)','Pa','Aproximación turbulenta E17')
add('CdA',f'={P("Qref")}/SQRT(2*{P("dpLow")}/{P("rho")})*{A("openingNozzle")}','m2','Equivalencia nominal turbulenta; exacta MSL Av/√2 en Comparacion')
add('diametro',f'=IF(OR({P("Cd")}="",{P("Cd")}<=0),"",SQRT(4*{A("CdA")}/(PI()*{P("Cd")})))','m','No se separan Cd y A sin dato externo')
add('vChorro',f'=IF(OR({A("diametro")}="",{A("Qmezcla")}=""),"",{A("Qmezcla")}/(PI()*{A("diametro")}^2/4))','m/s','Velocidad geométrica; diámetro real pendiente')
add('pBomba',f'=IF(OR({A("dpHose")}="",{A("dpGun")}="",{A("dpNozzle")}=""),"",{P("patm")}+{A("dpVenturi")}+{A("dpHose")}+{A("dpGun")}+{A("dpNozzle")})','Pa(abs)','Presupuesto condicional, no solución simultánea curva bomba')
add('dpHeater',f'={P("dpheat")}*({A("Q")}/{P("Qheatref")})*SQRT(({A("Q")}/{P("Qheatref")})^2+1e-8)','Pa','E04')
add('pSuccion',f'={P("patm")}-{A("dpHeater")}','Pa(abs)','Frontera del modelo; red real/filtro no representados')
add('dpBomba',f'=IF({A("pBomba")}="","",{A("pBomba")}-{A("pSuccion")})','Pa','Diferencia descarga-succión, no solo p manométrica')
add('QbombaCurva',f'=IF({A("dpBomba")}="","",MIN({P("Qref")},{P("Qref")}*MAX(0,(1+{P("slip")})*{A("omega")}-{P("slip")}*MAX(0,{A("dpBomba")})/{P("Pref")}),{P("eta")}*{P("Pel")}*{A("omega")}/MAX(1,{A("dpBomba")})))','m3/s','E06; cotejar Q, no curva comercial')
add('Ph',f'=IF({A("dpBomba")}="","",{A("Q") }*{A("dpBomba")})','W','E07; sin bypass en presupuesto')
add('H',f'=IF({A("dpBomba")}="","",{A("dpBomba")}/({P("rho")}*{P("g")}))','m','E07')
add('Peje',f'=IF(OR({A("Ph")}="",{P("etaPump")}="",{P("etaPump")}<=0),"",{A("Ph")}/{P("etaPump")})','W','E07 ηp eje→agua desconocida')

heat=sheet('Calentador_Modelo',['Variable','Valor','Unidad','Fuente / ecuación','Estado'])
heatrows=[('Potencia gas',f'={P("Pgas")}*{P("gasSetting")}','W','E05','Supuesto gas de entrada'),('Potencia útil al metal',f'=B2*{P("etaheat")}','W','E05','No potencia neta al agua'),('Caudal',f'={A("Q")}','m3/s','Editable Analitica_Modelo','Condicional'),('Flujo másico',f'=B4*{P("rho")}','kg/s','Continuidad','Calculado'),('T entrada',f'={P("Tin")}','°C','Usuario','13 °C'),('T agua estacionaria',f'=IF(B5<=0,"",(B3+B5*{P("cp")}*{P("Tin")}*(1+{P("Uamb")}/{P("UA")} )+{P("Uamb")}*{P("Tamb")})/(B5*{P("cp")}*(1+{P("Uamb")}/{P("UA")})+{P("Uamb")}))','°C','Solución dos balances E05 con derivadas cero','No aplica con quemador apagado/safety; extrapolación'),('T metal estacionaria',f'=IF(B7="","",B7+B5*{P("cp")}*(B7-B6)/{P("UA")})','°C','E05','Condicional'),('P al agua',f'=IF(B7="","",B5*{P("cp")}*(B7-B6))','W','m cp ΔT','Calculado'),('P ambiente',f'=IF(B8="","",{P("Uamb")}*(B8-{P("Tamb")}))','W','E05','Puede cambiar signo'),('Residual térmico', '=IF(B9="","",B3-B9-B10)','W','E05','Debe ~0 por ecuaciones'),('Viabilidad régimen', '=IF(B4<2/60000,"NO APLICA: umbral encendido",IF(B7>=60,"NO APLICA: corte térmico del modelo","Condicional: gas/sensores disponibles"))','—','Lógica del modelo; límites supuestos','No certificación'),('Compatibilidad 50 °C', '=IF(B7="","PENDIENTE",IF(B7>50,"Supera límite declarado","Bajo 50 °C; compatibilidad con agua caliente sigue PENDIENTE"))','—','Ficha contradictoria','Requiere fabricante')]
for row in heatrows:heat.append(row)

nps=sheet('NPSH_Modelo',['Variable','Valor','Unidad','Ecuación / fuente','Estado'])
nps.append(('p succión',f'={A("pSuccion")}','Pa(abs)','E04; frontera modelo','No presión Bogotá real'))
nps.append(('T succión','=IF(Calentador_Modelo!B12="Condicional: gas/sensores disponibles",Calentador_Modelo!B7,"")','°C','Estacionaria condicional','T simulada por separado'))
nps.append(('tau','=IF(OR(B3="",B3<0.01,B3>=373.946),"",1-(B3+273.15)/647.096)','1','IAPWS','Dominio líquido-saturación'))
nps.append(('p vapor','=IF(B4="","",22064000*EXP(647.096/(B3+273.15)*(-7.85951783*B4+1.84408259*B4^1.5-11.7866497*B4^3+22.6807411*B4^3.5-15.9618719*B4^4+1.80122502*B4^7.5)))','Pa(abs)','IAPWS SR1-86(1992)','Auxiliar; OMC no modela vapor'))
nps.append(('v brida',f'=IF(OR({P("Dsuccion")}="",{P("Dsuccion")}<=0),"",{A("Q")}/(PI()*{P("Dsuccion")}^2/4))','m/s','Continuidad','PENDIENTE diámetro succión'))
nps.append(('Margen estático',f'=IF(B5="","",(B2-B5)/({P("rho")}*{P("g")}))','m','(pabs-pv)/(rho g)','NO es NPSHA completa'))
nps.append(('NPSHA',f'=IF(OR(B6="",B7=""),"",B7+B6^2/(2*{P("g")}))','m','E19 en brida','PENDIENTE sin diámetro'))
nps.append(('NPSHr',f'=IF({P("NPSHr")}="","",{P("NPSHr")})','m','Fabricante','PENDIENTE'))
nps.append(('Margen bruto','=IF(OR(B8="",B9=""),"",B8-B9)','m','NPSHA-NPSHr; falta reserva','No validado'))
nps.append(('Estado','=IF(OR(B8="",B9=""),"PENDIENTE: geometría/NPSHr",IF(B10<=0,"NO CUMPLE","Pendiente reserva y validación fabricante"))','—','Cavitación no implementada','Simulación exitosa no demuestra ausencia'))

chem=sheet('Rama_Quimica',['Variable','Valor','Unidad','Ecuación / fuente','Estado'])
chem.append(('Presión manifold',f'=IF({A("pManifold")}="","",{A("pManifold")})','Pa(abs)','Frontera condicional Analitica_Modelo','PENDIENTE sin solución red o dato de ensayo'))
chem.append(('Nivel restante',f'=({P("Vtank")}-{P("drawnVolume")})/{P("Atank")}','m','E12','Sección constante; volumen provisional'))
chem.append(('A sifón',f'=PI()*{P("Dchem")}^2/4','m2','E02','12.7 mm interior confirmado'))
chem.append(('Elevación salida sifón',f'={P("lift")}+{P("Vtank")}/{P("Atank")}','m','ChemicalBranch.siphon.height_ab','Referencia fondo del tanque'))
chem.append(('Carga disponible',f'=IF(B2="","",{P("patm")}+{P("rho")}*{P("g")}*(B3-B5)-B2-{P("pcrack")})','Pa','Cabeza líquido - elevación - manifold - cracking rama','No incluye check HP del inyector: frontera en manifold'))
chem.append(('a resistencia laminar',f'=128*{P("mu")}*{P("Lchem")}/(PI()*{P("Dchem")}^4)','Pa/(m3/s)','Hagen–Poiseuille','Válida sólo régimen laminar/monofásico cebado'))
chem.append(('b resistencia válvulas',f'=IF(OR({P("uNC")}<=0,{P("needle")}<=0),"",({P("dpNC")}/{P("uNC")}^2+{P("dpCheck")}+{P("dpNeedle")}/{P("needle")}^2)/{P("qbranch")}^2)','Pa/(m3/s)^2','Tres válvulas nominales en serie','Aproximación fuera regularización'))
chem.append(('Q rama',f'=IF(OR({P("uNC")}=0,{P("needle")}=0),0,IF(OR(B6="",B8=""),"",IF(B6<=0,0,2*B6/(B7+SQRT(B7^2+4*B8*B6)))))','m3/s','Raíz positiva estable de aQ+bQ²=Δp disponible','Sin almacenamiento; no solución acoplada Venturi'))
chem.append(('v sifón','=IF(B9="","",B9/B4)','m/s','Q/A','Calculado'))
chem.append(('Re sifón',f'=IF(B10="","",{P("rho")}*ABS(B10)*{P("Dchem")}/{P("mu")})','1','E02','Verificar laminar'))
chem.append(('Δp sifón viscosa','=IF(B9="","",B7*B9)','Pa','Hagen–Poiseuille','No incluye elevación'))
chem.append(('Δp NC',f'=IF(B9="","",IF({P("uNC")}<=0,"",{P("dpNC")}*(B9/({P("qbranch")}*{P("uNC")}))^2))','Pa','E13 aproximada','Vacío y ratings inversos pendientes'))
chem.append(('Δp check',f'=IF(B9="","",{P("pcrack")}+{P("dpCheck")}*(B9/{P("qbranch")})^2)','Pa','E14 aproximada','Antirretorno ideal'))
chem.append(('Δp aguja',f'=IF(B9="","",IF({P("needle")}<=0,"",{P("dpNeedle")}*(B9/({P("qbranch")}*{P("needle")}))^2))','Pa','E15 aproximada','Sin curva real'))
chem.append(('p entrada NC',f'=IF(B12="","",{P("patm")}+{P("rho")}*{P("g")}*(B3-B5)-B12)','Pa(abs)','Balance de sifón con elevación','Estacionario condicional'))
chem.append(('p salida NC','=IF(OR(B16="",B13=""),"",B16-B13)','Pa(abs)','p entrada - Δp NC','Estacionario condicional'))
chem.append(('Validación cálculo','=IF(B11="","PENDIENTE",IF(B11>=2300,"NO APLICA: sifón no laminar",IF(B9=0,"Sin aspiración; presión atrapada requiere transitorio","Aproximación válida matemáticamente; requiere ensayo químico")))','—','Propiedades agua, cabeza disponible y válvulas supuestas','No validación física'))

# Snapshot comparisons evaluate the same constitutive laws on exported inputs.
cmp=sheet('Comparacion_Modelos',['Escenario','t s','Variable','Excel fórmula','Modelica / referencia','Unidad','Error absoluto','Error relativo %','Tipo de comparación','Causa / limitación','Umbral cercano a cero (unidad F)'])
def comp(name,t,label,formula,ref,unit,note,kind='Constitutiva sobre entradas simuladas'):
    row=cmp.max_row+1
    tol={'Pa':1,'Pa(abs)':1,'W':.001,'L/min':.00001,'°C':.000001}.get(unit,1e-9)
    cmp.append((name,t,label,'='+formula,'='+ref if isinstance(ref,str) else ref,unit,f'=IF(OR(D{row}="",E{row}=""),"",ABS(D{row}-E{row}))',f'=IF(OR(G{row}="",ABS(E{row})<K{row}),"",100*G{row}/ABS(E{row}))',kind,note+'; umbral relativo numérico de columna K, no tolerancia experimental',tol))
def X(n,t,k):return rawrefs.get((n,t,fullvar(k)))
for n,t,mode,r in snapshots:
    def x(k):return X(n,t,k)
    if not x('waterFlow_L_min'):continue
    qw=f'({x("waterFlow_L_min")}/60000)';qm=f'({x("outletFlow_L_min")}/60000)'
    dp=f'({x("pumpOutletPressure")}-{x("pumpInletPressure")})'
    comp(n,t,'P hidráulica bomba',f'({x("pumpFlow_L_min")}/60000)*MAX(0,{dp})',x('bomba.hydraulicPower'),'W','Incluye caudal total bomba y bypass; no potencia eje')
    comp(n,t,'Q curva equivalente',f'MIN({P("Qref")},{P("Qref")}*MAX(0,(1+{P("slip")})*{x("bomba.speed")}-{P("slip")}*MAX(0,{dp})/{P("Pref")}),{P("eta")}*{P("Pel")}*{x("bomba.speed")}/MAX(1,{dp}))*60000',x('pumpFlow_L_min'),'L/min','No curva fabricante; ω normalizada')
    comp(n,t,'Δp Venturi',f'{P("dploss")}*({qw}/{P("Qref")})*SQRT(({qw}/{P("Qref")})^2+.001^2)',x('injectorPressureDrop'),'Pa','E09; coeficientes sin calibrar')
    comp(n,t,'p garganta',f'MAX({P("pfloor")},{x("injectorOutletPressure")}-{P("recovery")}*({qw}/{P("Qref")})^2)',x('throatPressure'),'Pa(abs)','E09; no geometría AQB')
    xx=f'(({x("chemicalManifoldPressure")}-{x("throatPressure")}-{P("pcrack")})/{P("dpchem")})'
    root=f'IF({xx}<=0,0,IF({xx}<.001,SQRT(.001)*(2.5*({xx}/.001)^2-1.5*({xx}/.001)^3),SQRT({xx})))'
    comp(n,t,'Q aspirado Venturi',f'{P("qchemref")}*{root}*MAX(0,{P("rho")}*{qw})/SQRT(({P("rho")}*{qw})^2+{P("mreg")}^2)*60000',f'{x("injector.mChemical")}*60000/{P("rho")}','L/min','No caudal impuesto; diferencial y resistencias supuestas')
    comp(n,t,'Masa inyector',f'{x("injector.mWater")}+{x("injector.mChemical")}',x('injector.mMixture'),'kg/s','E11; verifica unión')
    mixsum='+'.join(x(f'injectedProductFlow[{j}]') for j in [1,2,3])
    comp(n,t,'Fracción producto inyector',f'({mixsum})/MAX(1E-8,{x("injector.mMixture")})',x('chemicalFractionAtInjector'),'1','E18; equivalente másico, no concentración activa')
    comp(n,t,'Concentración trazador salida','+'.join(x(f'outletConcentration[{j}]') for j in [1,2,3]),x('outletChemicalFraction'),'1','Suma de trazadores transitorios; no confundir con mezcla instantánea')
    comp(n,t,'p succión',f'{P("patm")}-{P("dpheat")}*(({x("pumpFlow_L_min")}/60000)/{P("Qheatref")})*SQRT((({x("pumpFlow_L_min")}/60000)/{P("Qheatref")})^2+1e-8)',x('pumpInletPressure'),'Pa(abs)','No filtro/red suministro')
    v=f'({qm}/(PI()*{P("Dhose")}^2/4))';reyn=f'({P("rho")}*ABS({v})*{P("Dhose")}/{P("mu")})'
    comp(n,t,'Velocidad manguera',v,x('hose.flowModel.vs[1]'),'m/s','Q mezcla / A; geometría supuesta, no medida')
    ffr=f'IF({reyn}<2300,64/MAX(1E-30,{reyn}),.25/(LOG10({P("epshose")}/(3.7*{P("Dhose")})+5.74/MAX(1E-30,{reyn})^.9)^2))'
    comp(n,t,'Δp manguera Darcy',f'IF({qm}=0,0,{ffr}*{P("Lhose")}/{P("Dhose")}*{P("rho")}*{v}^2/2)',x('hosePressureDrop'),'Pa','MSL from_dp=true inversión Colebrook; Excel Swamee–Jain; transición distinta')
    av=f'({P("rho")}*{P("Qref")}/(SQRT({P("rho")})*({P("dpLow")}/({P("dpLow")}^2+1)^.25)))'
    comp(n,t,'Caudal boquilla nominal Av',f'{x("gun.nozzle.opening")}*{av}*SQRT({P("rho")})*SQRT(MAX(0,{x("nozzlePressureDrop")}))*60000/{P("rho")}',x('outletFlow_L_min'),'L/min','Av calculado en Excel con OpPoint/regRoot y dp_small=1 Pa; sqrt flujo sólo fuera regularización, no Av exportado')
    comp(n,t,'P al agua transitoria',f'{P("rho")}*({x("pumpFlow_L_min")}/60000)*{P("cp")}*({x("inletHeater.T")}-{P("Tin")}-273.15)',x('inletHeater.heatToWater'),'W','Diferencia por acumulación agua; NO error si transitorio','Balance sin término de almacenamiento para comparar')
    comp(n,t,'Balance transitorio agua calentador',f'{P("rho")}*{P("Vwater")}*{P("cp")}*{x("der(inletHeater.T)")}',f'{x("inletHeater.heatToWater")}-{P("rho")}*({x("pumpFlow_L_min")}/60000)*{P("cp")}*({x("inletHeater.T")}-{P("Tin")}-273.15)','W','Inventario .3 L del modelo y flujo directo','Conservación')
    comp(n,t,'Balance transitorio metal',f'{P("Cm")}*{x("der(inletHeater.metalTemperature)")}',f'{x("heatingPower")}-{x("inletHeater.heatToWater")}-{x("inletHeater.ambientLoss")}','W','Capacidad metal 1500 J/K del modelo','Conservación')
    comp(n,t,'Residual energía calentador',x('inletHeater.energyResidual'),0,'W','Balance dos estados; error relativo omitido referencia cero','Conservación')
    for b in ['prewash','soap','wax']:
        qc=f'({x(b+"m_flow")})' if x(b+'m_flow') else None
        mnom=f'{P("rho")}*'+(f'{P("qbranch")}' if b!='wax' else '(.10/60000)')
        z=f'({x(b+".valveDifferentialPressure")}/{P("dpNC")})'
        comp(n,t,b+' NC flujo',f'{mnom}*{x(b+".solenoid.u")}*{z}/({z}^2+.001^2)^.25',x(b+'.solenoid.m_flow'),'kg/s','E13 exacta; NC cerrada sin fuga')
        z=f'({x(b+".needle.dp")}/{P("dpNeedle")})'
        comp(n,t,b+' aguja flujo',f'{mnom}*{P("needle")}*{z}/({z}^2+.001^2)^.25',x(b+'.needle.m_flow'),'kg/s','E15; no geometría física')
        comp(n,t,b+' masa almacenamiento',f'{x(b+".tankDraw")}-{x(b+".m_flow")}-{x(b+".beforeCheck.storageRate")}-{x(b+".afterCheck.storageRate")}',0,'kg/s','Caudal tanque ≠ aguja durante transitorio','Conservación')
    for key in ['mainMassResidual','chemicalMassResidual','injectorMassResidual']:comp(n,t,key,x(key),0,'kg/s','Masa conservada incluyendo almacenamiento','Conservación')
for name in ['HacebFull','HacebPower']:
    rows,_,_=runs[name];t=rows[-1]['time']
    m=rawrefs[(name,t,'heater.m_flow')];power=rawrefs[(name,t,'heater.heaterPower')]
    numerator=f'({power}+{m}*{P("cp")}*{P("Tin")}*(1+{P("Uamb")}/{P("UA")})+{P("Uamb")}*{P("Tamb")})'
    denominator=f'({m}*{P("cp")}*(1+{P("Uamb")}/{P("UA")})+{P("Uamb")})'
    comp(name,t,'T calentador equilibrio',f'{numerator}/{denominator}',rawrefs[(name,t,'heater.T')]+'-273.15','°C','E05 derivadas cero; comparar sólo banco aproximado a equilibrio','Equilibrio analítico frente final transitorio')

stations=sheet('Estaciones_Modelica',['Escenario','t s','Estación','p absoluta Pa','p manométrica bar','T °C','Q L/min','rho kg/m3','Di m','v m/s','Fuente / estado'])
for n,t,mode,r in snapshots:
    for lab,pressure,temp,flow,dia in [('Fuente',None,None,'pumpFlow_L_min',None),('Salida calentador / succión','pumpInletPressure','inletHeater.T','pumpFlow_L_min','Dsuccion'),('Descarga / acumulador','pumpOutletPressure','waterTemperature','pumpFlow_L_min',None),('Entrada Venturi','injectorInletPressure','waterTemperature','waterFlow_L_min',None),('Garganta','throatPressure',None,'waterFlow_L_min',None),('Salida Venturi','injectorOutletPressure',None,'outletFlow_L_min',None),('Entrada pistola','gunInletPressure','outletTemperature','outletFlow_L_min','Dhose'),('Entrada boquilla',None,'outletTemperature','outletFlow_L_min',None),('Descarga ambiente',None,'outletTemperature','outletFlow_L_min',None),('Manifold químico','chemicalManifoldPressure',None,None,None)]:
        rr=stations.max_row+1
        pressureformula=f'={X(n,t,pressure)}' if pressure else (f'={P("patm")}' if lab in ['Fuente','Descarga ambiente'] else f'={X(n,t,"gunInletPressure")}-{X(n,t,"gun.pistola.dp")}' if lab=='Entrada boquilla' else None)
        tf=f'={X(n,t,temp)}-273.15' if temp else f'={P("Tin")}' if lab=='Fuente' else None
        qf=f'={X(n,t,flow)}' if flow else None
        df=f'=IF({P(dia)}="","",{P(dia)})' if dia else None
        stations.append((n,t,lab,pressureformula,f'=IF(D{rr}="","",(D{rr}-{P("patm")})/100000)',tf,qf,f'={P("rho")}',df,f'=IF(OR(G{rr}="",I{rr}="",I{rr}<=0),"",G{rr}/60000/(PI()*I{rr}^2/4))','Instantánea; T/Q sin datos se dejan vacíos; no temperatura exacta de garganta'))

sup=sheet('Supuestos',['ID','Valor / hipótesis','Tipo','Justificación','Efecto','Cómo cerrar'])
for d in defs:
    if 'supuesto' in d[3].lower() or d[5]=='PENDIENTE':sup.append((d[0],d[1] if d[1] is not None else 'PENDIENTE',d[3],d[4],d[5],'Ficha exacta o medición'))
for row in [('Agua equivalente','Productos rho/mu agua','Hipótesis','Modelo actual','Dosis y pérdida química no validadas','Ficha formulación prelavado/jabón/cera'),('Tapa','Con respiradero','Decisión usuario','Sólo variante vigente','Sin depresión de espacio gas; posible obstrucción no modelada','Definir respiradero compatible'),('Cebado','Manguera llena siempre','Hipótesis','Modelo monofásico','No predice arranque de aspiración ni burbujas','Ensayo cebado'),('Ideal check','Sin fuga inversa','Hipótesis','Dos niveles de protección','Presión NC depende fallos; no rating seleccionado','Ficha y ensayo fallo seguro'),('Mezcla','Trazador pasivo separado de Medium','Hipótesis','Tres escalares','No reacciones ni viscosidad real','Medir densidad, viscosidad y concentración'),('NPSH','Sin módulo cavitación OMC','Limitación','ConstantPropertyLiquidWater','Simulación no verifica NPSH','Medir succión; solicitar NPSHr')]:sup.append(row)

sens=sheet('Sensibilidad',['Caso','Q agua L/min','T succión °C','Patm Pa(abs)','p suministro bar(g)','L m','D m','epsilon m','p salida Venturi Pa(abs)','dpFiltro Pa','dpCalentador Pa','Cd supuesto','d boquilla supuesto m','rho kg/m3','mu Pa.s','v m/s','Re','f Darcy','dp manguera Pa','p garganta Pa(abs)','p vapor Pa(abs)','Margen estático succión m','dp boquilla Pa','Estado'])
cases=[('Base modelo',7,34,101325,0,8,.008,2.5e-5,250000,0,None,None,None),('Diseño manguera 4m',7,34,101325,0,4,.008,2.5e-5,250000,0,None,None,None),('Bogotá supuesto plantilla',7,34,74000,0,8,.008,2.5e-5,250000,0,None,None,None),('Caudal 5',5,34,101325,0,8,.008,2.5e-5,250000,0,None,None,None),('Caudal 9 fuera ficha',9,34,101325,0,8,.008,2.5e-5,250000,0,None,None,None),('Agua fría',7,13,101325,0,8,.008,2.5e-5,250000,0,None,None,None),('Agua 50',7,50,101325,0,8,.008,2.5e-5,250000,0,None,None,None),('Diam 6mm',7,34,101325,0,8,.006,2.5e-5,250000,0,None,None,None),('Diam 10mm',7,34,101325,0,8,.010,2.5e-5,250000,0,None,None,None),('Rugosidad 5um',7,34,101325,0,8,.008,5e-6,250000,0,None,None,None),('Mayor contrapresión',7,34,101325,0,8,.008,2.5e-5,400000,0,None,None,None),('Red hipotética 1bar',7,34,101325,1,8,.008,2.5e-5,250000,0,None,None,None),('Filtro hipotético 10kPa',7,34,101325,0,8,.008,2.5e-5,250000,10000,None,None,None),('Calentador hipotético 40kPa',7,34,101325,0,8,.008,2.5e-5,250000,0,40000,None,None),('Boquilla diámetro pendiente',7,34,101325,0,8,.008,2.5e-5,250000,0,None,None,None)]
for case in cases:
    rr=sens.max_row+1;vals=list(case)
    if vals[10] is None:vals[10]=f'={P("dpheat")}*(B{rr}/60000/{P("Qheatref")})^2'
    sens.append(vals+[f'={P("rho")}',f'={P("mu")}',f'=B{rr}/60000/(PI()*G{rr}^2/4)',f'=N{rr}*P{rr}*G{rr}/O{rr}',f'=IF(Q{rr}<2300,64/Q{rr},IF(Q{rr}<4000,"",.25/(LOG10(H{rr}/(3.7*G{rr})+5.74/Q{rr}^.9)^2)))',f'=IF(R{rr}="","",R{rr}*F{rr}/G{rr}*N{rr}*P{rr}^2/2)',f'=MAX({P("pfloor")},I{rr}-{P("recovery")}*(B{rr}/60000/{P("Qref")})^2)',f'=IF(OR(C{rr}<.01,C{rr}>=373.946),"",22064000*EXP(647.096/(C{rr}+273.15)*(-7.85951783*(1-(C{rr}+273.15)/647.096)+1.84408259*(1-(C{rr}+273.15)/647.096)^1.5-11.7866497*(1-(C{rr}+273.15)/647.096)^3+22.6807411*(1-(C{rr}+273.15)/647.096)^3.5-15.9618719*(1-(C{rr}+273.15)/647.096)^4+1.80122502*(1-(C{rr}+273.15)/647.096)^7.5)))',f'=IF(U{rr}="","",(D{rr}+E{rr}*100000-J{rr}-K{rr}-U{rr})/(N{rr}*{P("g")}))',f'=IF(OR(L{rr}="",L{rr}<=0,M{rr}="",M{rr}<=0),"",N{rr}/2*(B{rr}/60000/(L{rr}*PI()*M{rr}^2/4))^2)','Ensayo condicional hipotético; no punto real; agua MSL constante'])
chart=ScatterChart();chart.title='Longitud y pérdida: casos condicionales';chart.x_axis.title='Longitud m';chart.y_axis.title='Δp Pa'
chart.series.append(Series(Reference(sens,min_col=19,min_row=2,max_row=3),Reference(sens,min_col=6,min_row=2,max_row=3),title='8 m / 4 m'))
sens.add_chart(chart,'Z2')
for j,title in enumerate(['d/d0 boquilla','Q/Q0','dp/dp0','Interpretación'],26):sens.cell(20,j,title)
for rr,ratio in enumerate([.8,1,1.2],21):
    sens.cell(rr,26,ratio);sens.cell(rr,27,1);sens.cell(rr,28,f'=AA{rr}^2/Z{rr}^4');sens.cell(rr,29,'Sensibilidad relativa orificio, mismo Cd; no identifica d0 ni presión real')

verify=sheet('Integridad',['Criterio','Evidencia','Estado','Acción pendiente'])
for row in [('Masa nodos','Residuos OMC + comparaciones almacenamiento','Verificado numéricamente en escenarios','Ensayo físico'),('Energía calentador','energyResidual; dos estados','Verificado numéricamente','Calibrar UA, eficiencia, volumen'),('Energía sistema','Sin trabajo compresión, calor bomba/disipación; mezcla regularizada','Parcial','Balance energético completo'),('Presión conexiones','Puertos conectados comparten p; fuente/hose/nozzle exportados','Verificado por ecuaciones','Sensores prototipo'),('Curva bomba-boquilla','Q depende carga equivalente; no curva medida','Aproximado','Curva fabricante/ensayo'),('Ratings componentes','No disponibles','PENDIENTE','No elegir NC por presión normal solamente'),('NPSH','pv auxiliar IAPWS; sin diámetro brida/NPSHr','PENDIENTE','Ficha + medición succión'),('Compatibilidad térmica','50 °C vs no apta caliente','PENDIENTE CRÍTICO','Confirmación referencia real'),('Aspiración AQB','LP aspira en modelo; HP no','Aproximado','Referencia exacta y mapa aspiración'),('Cierre','Autostop, elasticidad, alivio; presión atrapada ~117 bar','Simulado; mecanismo real pendiente','Transitorio real y golpe ariete'),('Materiales químicos','PP propuesto; formulaciones desconocidas','PENDIENTE','HDS, sellos, temperatura, vacío'),('Agua/mezcla','Propiedades agua constantes; químicos trazadores','Aproximado','rho,mu productos y mezcla'),('Resultado experimental','Ninguno','PENDIENTE','Plan de pruebas'),('Filtro','Ausente en modelo','PENDIENTE','Seleccionar y medir limpio/cargado')]:verify.append(row)

for title,url,note in [('IAPWS SR1-86(1992)','https://iapws.org/technical-guidance/release/Supp-sat','Coeficientes presión vapor; consulta 2026-10-09'),('AQB producto','https://www.aqbasociados.com.co/producto/inyectores-quimicos-venturi/','Sin curvas cuantitativas de referencia exacta en página consultada'),('HACEB catálogo oficial','https://www.haceb.com/calentadores/de-paso','No identifica equipo usado suministrado; no sustituir ficha exacta'),('MSL fuente local','C:/Users/Mateo/AppData/Roaming/.openmodelica/libraries/Modelica 4.1.0+maint.om','Pipes.mo / Valves.mo / Media/Water/package.mo inspeccionados')]:wb['Fuentes'].append((f'S-{wb["Fuentes"].max_row}',title,note,url,'No certifica prototipo'))

for s in list(wb)[len(inspection):]:
    if s.title=='Parametros_Modelica':
        for row in s.iter_rows(min_row=2):
            # MSL descriptions often begin '= true ...'; they are source text,
            # never executable Excel expressions.
            row[6].data_type='s'
    s.auto_filter.ref=s.dimensions
    s.row_dimensions[1].height=32
    for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='17365D');c.alignment=Alignment(wrap_text=True,vertical='center')
    for col in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(col)].width=25 if col<5 else 42
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.data_type=='f':c.fill=PatternFill('solid',fgColor='E2F0D9');c.number_format='0.000000'
            elif isinstance(c.value,(int,float)):c.fill=PatternFill('solid',fgColor='DDEBF7');c.number_format='0.000000'
            elif c.value is None or 'PENDIENTE' in str(c.value):c.fill=PatternFill('solid',fgColor='FFF2CC')
    s.sheet_view.zoomScale=85
for s,rng in [('Entradas_Modelo',f'B2:B{inp.max_row}'),('Sensibilidad',f'B2:M{sens.max_row}'),('Analitica_Modelo',f'B2:B{a.max_row}')]:
    dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);dv.errorTitle='Dato inválido';dv.error='Ingrese valor no negativo; deje vacío si no está identificado.';dv.showErrorMessage=True
    # Pressure/temp rows can be below zero; targeted positive inputs only.
    if s=='Entradas_Modelo':
        wb[s].add_data_validation(dv)
        for k in ['rho','mu','cp','g','Qref','Lhose','Dhose','epshose','Dchem','Dsuccion','NPSHr','Cd','Pel']:dv.add(f'B{R[k]}')
    elif s=='Sensibilidad':wb[s].add_data_validation(dv);dv.add(f'B2:B{sens.max_row}');dv.add(f'F2:H{sens.max_row}')
    else:wb[s].add_data_validation(dv);dv.add(f'B{ar["Q"]}');dv.add(f'B{ar["Qquim"]}')
for k in ['eta','etaheat','gasSetting','needle','slip']:
    dv=DataValidation(type='decimal',operator='between',formula1=0.000001 if k in ['eta','etaheat'] else 0,formula2=1,allow_blank=True);dv.showErrorMessage=True;inp.add_data_validation(dv);dv.add(f'B{R[k]}')
positive=DataValidation(type='decimal',operator='greaterThan',formula1=0,allow_blank=True)
positive.showErrorMessage=True;positive.error='Este parámetro debe ser positivo; dejar vacío si falta identificación.'
inp.add_data_validation(positive)
for k in ['rho','mu','cp','g','Qref','Qheatref','Dhose','Dchem','Dsuccion','Atank','qbranch','UA','dpchem','dpNC','dpCheck','dpNeedle','dpLow','dpHigh','Cd']:positive.add(f'B{R[k]}')
for k in ['uNC']:
    dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=1,allow_blank=True);dv.showErrorMessage=True;inp.add_data_validation(dv);dv.add(f'B{R[k]}')
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True)
wb.save(OUT)

# Reports preserve complete local hierarchy/parameters in linked workbook rather than dumping MSL internals.
issues=[
('Crítica','Referencia bomba inconsistente','Excel 2100 W/110/150 bar; modelo 1800 W/100 supuesto/130 bar. Ninguna curva oficial identificada. Conservar ambos; fotografiar placa y obtener manual.'),
('Crítica','Red y filtro ausentes','Fuente Boundary_pT a 101325 Pa(abs), no red presurizada Bogotá. Sin filtro ni tuberías de succión. Pérdida calentador genera succión subatmosférica; NPSH real no cerrado.'),
('Crítica','Venturi no caracterizado','Pérdida y recuperación 180 kPa a 7 L/min son coeficientes supuestos. Sin geometría, curva AQB, cavitación ni propiedades del producto.'),
('Crítica','Compatibilidad agua caliente','Usuario proporciona 50 °C máximo y no apta para caliente. El assert del modelo no es autorización del fabricante.'),
('Alta','Protección química ideal','Dos checks ideales; ensayo de presión impuesta de 130 bar no es simulación de fallo del inyector. Una NC puede recibir diferencial inverso de ~130 bar si falla check y se presuriza manifold.'),
('Alta','NPSH sin cerrar','Sin NPSHr ni diámetro brida succión; presión vapor sólo postprocesada. pFloor=20 kPa y assert p>10 kPa no son criterios de cavitación dependientes de T.'),
('Alta','Cierre simplificado','Autostop basado gatillo; motor .2 s y volúmenes elásticos lumped. Alivio equivalente 117 bar, sin histéresis. Presión atrapada, sin golpe de ariete, pulsación ni certificación alivio real.'),
('Alta','Potencia bomba y eficiencia','η=.70 es conversión eléctrica-hidráulica para limitar potencia. No ηp de eje; velocidad normalizada no rpm. No calor motor/bomba ni cavitación.'),
('Media','Calentador no identificado','12 kW asumidos gas entrada, η=.85, UA=800 W/K, Uamb=2 W/K, Cm=1500 J/K, agua .3 L. Referencia exacta y mínimos de alimentación pendientes.'),
('Media','Mezcla aproximada','Productos son escalares pasivos externos a Medium. Igual rho/mu/cp agua; depósito a temperatura constante, sin cebado/aire/reacción.'),
('Media','Entalpía inyector regularizada','mreg añade transporte entálpico ficticio regularizador; diferencia respecto mezcla estricta proporcional a mreg y diferencia de entalpías. Masa principal sí exacta.'),
('Media','Caudal químico mal descrito','chemicalFlow_kg_s suma caudales de agujas al manifold; docstring dice extraído de tanques. Usar tankDraw para extracción; almacenamiento explica diferencia.'),
('Media','Trazador en reversión','El balance de concentración presupone aportes de producto puros y usa suma firmada de caudales de ramas. No transporta composición real de reflujo en ramas/compliance; revisar si hay reversión.'),
('Media','Geometría hose','8 m modelo vs 4 m propuesta. Di 8 mm y rugosidad 25 μm supuestos; fittings K no identificados.'),
('Menor','Documentación obsoleta','Comentario del paquete enlaza verification/autolavado.original.mo ya eliminado; rama llama .0127 m asumido aunque usuario confirmó diámetro interior.'),
('Menor','Manual en cero','gasValveOpen/burnerLit pueden permanecer verdaderos con gasSetting=0 aunque gasPower=0. Revisar coherencia lógica, sin alterar modelo durante auditoría.'),
('Corrección Excel','Referencia pérdida pistola','Balances apuntaban B8 (modelo/nombre) en vez B9 (Δp). Corregido en copia.'),
('Corrección Excel','Pendientes y transición','Blancos no son cero: sin salida calentador no asumir agua fría; SUMIF no aceptar tramos incompletos; transición sin correlación elegida no calculada.')]
report=['# Auditoría del modelo hidráulico','',f'Fecha generación UTC: {now}. Modelo SHA256: `{sha}`. Excel original SHA256: `{original_sha}`.','',
'El libro conserva las 14 hojas originales y todos los datos comerciales suministrados. Las modificaciones de fórmulas y la frontera de 13 °C están registradas. El modelo `.mo` no se modifica. Los números simulados no son mediciones del prototipo.','',
'## Arquitectura y fronteras','',
'`SelfServiceCarWashTest` conecta `TestSequence` con `plant`. Agua: `fuente.ports[1] → inletHeater → bomba → accumulator → injector.water → injector.mixture → hose → gun.pistola → gun.nozzle → salida`. Bypass: accumulator → unloader → fuente.ports[2]. Tres ramas: tanque ventilado → StaticPipe sifón → NC 12 V → volumen beforeCheck → check → volumen afterCheck → aguja → manifold común → check HP integrado en injector. Puertos FluidPort transportan p absoluta, m_flow firmado y entalpía stream; las conexiones conservan masa y comparten p. Dirección declarada no garantiza ausencia de reflujo: TwoPort/StaticPipe/aguja/NC abierta permiten reversión, checks ideales la restringen.','',
'Fuente y salida: 101325 Pa(abs), 286.15 K. Tanques químicos con respiradero al mismo ambiente. No existe filtro de agua explícito ni acometida de Bogotá. Calentador separado aguas arriba, ajuste manual 1 y sin control automático de temperatura. Estado térmico inicial agua/metal 13 °C; acumulador y colectores p0=patm, T0=13 °C, motor y trazadores iniciales cero.','',
'## Jerarquía completa inspeccionada','',
'`verification/model-instance.json` conserva la instancia completa de OMC. Hojas Componentes, Conexiones, Declaraciones y Parametros_Modelica contienen todas las instancias, conexiones, declaraciones de la fuente y parámetros resueltos o expresiones pendientes. Se incluyen clases MSL internas, control y bancos de ensayo. No se interpretan valores start de variables dinámicas como condiciones estacionarias.','',
'| Ruta | Clase | Función / estado |','|---|---|---|']
for p,t,comment,src in components:
    report.append(f'| `{p}` | `{t}` | {comment.replace("|","/")} |')
report+=['','## Ecuaciones físicas y trazabilidad','','| ID | Subsistema | Relación | Fuente / hipótesis |','|---|---|---|---|']
for eid,sub,formula,cl,hyp in eqs:report.append(f'| {eid} | {sub} | {formula.replace("|","abs")} | {cite(cl)}; {hyp} |')
report+=['','MSL inspeccionada: Media/Water/package.mo:50–54 (rho=995.586 kg/m³, mu=.001 Pa·s, cp=4184 J/kgK); Fluid/Pipes.mo:StaticPipe (balances sin almacenamiento); PartialGenericPipeFlow:1123 (`from_dp=true` en StaticPipe); WallFriction.Detailed:2879 y funciones internas Colebrook/Swamee–Jain; Fluid/Valves.mo:ValveIncompressible:35–64 (Av y regularizaciones). La inversión m(dp) usa Colebrook; la fórmula directa dp(m) usa Swamee–Jain. Hay interpolación cúbica en transición, regularización de caudal y altura. Excel usa Darcy laminar/Swamee–Jain turbulento y deja la transición original pendiente. Diferencias pequeñas de pérdida no son un fallo de conservación.','',
'La unidad de conversión PaPerPsi está declarada Pa aunque representa Pa/psi; factores 60000 y 6894.757 son numéricos. p absoluta y diferencias de p se separan en la memoria. QΔp está en W, m cp ΔT en W y Δp/(rho g) en m. `checkModel` con `--unitChecking` no emitió advertencias; eso no demuestra que cada factor de conversión esté semánticamente tipado ni que las ecuaciones físicas sean válidas.','',
'## Hallazgos y acciones','','| Prioridad | Hallazgo | Consecuencia / acción |','|---|---|---|']
for pr,head,body in issues:report.append(f'| {pr} | {head} | {body} |')
report+=['','## Simulación y comparación','',
'DASSL, tolerancia 1e-6; secuencia ciclo 0–64 s con intervalo .02 s; dosificación ventilada 0–120 s con .02 s; protección/fallo 0–15 s con .01 s. HeatingTest 0–30 s .02 s. Bancos calentador: plena potencia/sobretemperatura/potencia insuficiente 120 s, fallo llama/ventilación/bajo caudal 40 s, .1 s. Escenarios, fechas y SHA quedan en CSV y libro. Bancos de protección imponen 60 kPa(abs) antes de 5 s y 130 bar(g) después; NC cerrada desde 2 s; failedCheck cambia sólo en banco de fallo.','',
'Las muestras de 9,18,33,48 y58 s son instantáneas HP/prelavado/jabón/cera/cierre. El calentador y trazadores pueden seguir evolucionando. Dosificación a 120 s tampoco es un estacionario estricto: el tanque pierde nivel. Banco HacebFull al final se aproxima a equilibrio térmico; confirmar derivadas antes de reportar estacionario. Comparacion_Modelos verifica leyes constitutivas usando entradas exportadas y lo rotula expresamente; no se presenta como predicción independiente. Analitica_Modelo deja planteado presupuesto de red para Q, dosis, contrapresión y posiciones editables; no resuelve automáticamente las ecuaciones acopladas con bomba y ramas.','',
'NPSH y velocidad de garganta no se comparan con OMC porque no existen allí ni está identificada geometría; quedan pendientes. La velocidad en manguera es derivada Q/A, no una medición. CdA exacto del modelo se obtiene de Av/√2; no se deducen Cd y diámetro por separado. Peje y rpm pendientes. Concentración es trazador de producto equivalente, no concentración activa comercial.','',
'## Fuentes externas y límites de verificación','',
'[AQB, inyector Venturi](https://www.aqbasociados.com.co/producto/inyectores-quimicos-venturi/) describe el principio, pero el contenido consultado no entrega la curva cuantitativa de una referencia exacta. [IAPWS SR1-86(1992)](https://iapws.org/technical-guidance/release/Supp-sat) respalda la correlación de presión de vapor usada en Excel. [Catálogo Haceb](https://www.haceb.com/calentadores/de-paso) no identifica el equipo usado proporcionado. No se obtuvo manual oficial TRENT exacto; datos comerciales siguen sin confirmar. Consultas: 2026-10-09.','',
'No se inventan parámetros del filtro ni geometría AQB. Validación experimental pendiente: presión/caudal simultáneos de bomba, pérdidas limpias/cargadas, temperatura máxima real, aspiración por contrapresión, concentración y compatibilidad química; ensayo de cierre con protecciones adecuadas. La coincidencia Excel–OMC sólo demuestra consistencia entre ecuaciones.','',
'## Reproducción','',
'1. Ejecutar `verification/verify.ps1` para inyección y después `verification/verify.ps1 -Script verification/haceb.mos` (secuencial: comparten log).\n2. Ejecutar `omc verification/audit.mos`.\n3. Ejecutar `python verification/build_memory.py`.\n4. Recalcular y guardar la copia con Excel mediante `verification/recalculate_memory.ps1`; ejecutar `python verification/validate_memory.py`.\n5. No modificar ni sustituir el Excel original. Revisar los parámetros pendientes antes de usar el presupuesto como diseño.']
(ROOT/'Informe_Auditoria_Modelo.md').write_text('\n'.join(report),encoding='utf-8')
pending=['# Parámetros pendientes','', 'Los datos simulados o supuestos no sustituyen las fichas del equipo. Prioridades para cerrar el diseño:','', '| Prioridad | Dato | Motivo / obtención |','|---|---|---|',
'| Crítica | Placa/manual exacto TRENT, curva Q–Δp y protección | Resolver 1800/2100 W y 130/150 bar; verificar 7 L/min simultáneos, presión admisible y cierre. |',
'| Crítica | Autorización de agua precalentada | Ficha contradictoria: 50 °C máximo y no apta caliente; identificar temperatura admisible continua. |',
'| Crítica | AQB referencia, orificio/diámetros, curva aspiración y límite contrapresión | Sin curva no seleccionar punto de dosificación ni dar por válido caudal químico. |',
'| Crítica | Ratings HP check inyector, rama, manifold, aguja, NC y manguera química | Presión normal de vacío no cubre fallo de aislamiento. Confirmar fugas inversas, Δp inversa, vacío y acción directa. |',
'| Alta | Presión dinámica de acometida y presión barométrica local | Medición con caudal, mínimos y variación temporal; 74 kPa es supuesto plantilla, OMC usa101325 Pa. |',
'| Alta | NPSHr y reserva, diámetro/cota brida succión, pérdidas de acometida | Sin ellos NPSHA completa y margen no evaluables. |',
'| Alta | Filtro referencia, curva Δp–Q limpio/cargado | No existe filtro en modelo; no sustituir pérdida desconocida por cero. |',
'| Alta | HACEB placa, 12 kW gas/útil, mínimos de presión/caudal, curva Δp–Q | η, inventario, UA, inercia, encendido y temperatura de corte son supuestos. |',
'| Alta | Longitud/diámetro/rugosidad/cotas y K de todas las líneas | 4 m diseño frente8 m modelo; Di HP8mm supuesto; químicos Di12.7mm confirmado, L2m/lift.5m pendientes. |',
'| Alta | Prelavado/jabón/cera: HDS, composición, rho(T), mu(T), dilución | Compatibilidad PP, sellos NC/check/aguja y mangueras; densidad/viscosidad agua no verificadas. |',
'| Alta | Tanques volumen/sección/nivel min-max; respiradero y cebado | Sólo variante ventilada; caída de nivel cambia aspiración. |',
'| Alta | Pistola/boquilla exactas, curva, Cd o área, patrón de chorro | Av equivalente no define diámetro ni Cd individual; comprobar presión/temperatura máximas. |',
'| Media | ηp, ηmotor, rpm, disipación/elasticidad | No confundir eficiencia global de modelo ni omega normalizada con valores medidos. |',
'| Media | Ensayos estacionarios y transitorios | Registrar Q,p,T/dosis simultáneos, tiempos de purga, cierre, presiones atrapadas y ausencia de contaminación cruzada. |']
(ROOT/'Parametros_Pendientes.md').write_text('\n'.join(pending),encoding='utf-8')
log=['# Registro de cambios','',f'Original conservado: `{ORIGINAL.name}`; SHA256 `{original_sha}`. Copia: `{OUT.name}`. Modelo sin cambios; SHA256 `{sha}`.','', '## Inspección previa','']
for n,st in inspection.items():log.append(f'- {n}: {st["rows"]} filas × {st["cols"]} columnas, {st["formulas"]} fórmulas, {st["values"]} celdas con datos.')
log+=['','## Modificaciones en hojas originales','','| Hoja!celda | Original | Nuevo | Motivo |','|---|---|---|---|']
for s,c,old,new,reason in changes:log.append(f'| {s}!{c} | {str(old).replace("|","/")} | {str(new).replace("|","/")} | {reason} |')
log+=['','## Hojas añadidas','']
for s in list(wb)[len(inspection):]:log.append(f'- {s.title}: {s.max_row-1} registros; fórmulas verdes, números editables/datos azules y pendientes amarillos. Clasificación y fuentes explícitas.')
log+=['','## Verificación','','Revisión estructural, recálculo nativo Excel y controles numéricos quedan en verification/results/memory-validation.json al ejecutar validate_memory.py. Exportaciones OMC con fecha/escenario/unidad y SHA de fuente; no datos experimentales. No se modifica el modelo para lograr coincidencia. El CSV conserva historias completas; el Excel almacena instantáneas para tamaño manejable.','',f'Generación UTC: {now}.']
(ROOT/'Registro_Cambios.md').write_text('\n'.join(log),encoding='utf-8')
(ROOT/'verification/results/memory-build.json').write_text(json.dumps({'date':now,'source_sha256':sha,'original_sha256':original_sha,'original_sheets':inspection,'new_sheets':wb.sheetnames[len(inspection):],'csv_rows':len(export),'comparisons':cmp.max_row-1,'changes':len(changes),'output':str(OUT)},indent=2),encoding='utf-8')
print(json.dumps({'workbook':str(OUT),'sheets':len(wb.sheetnames),'comparisons':cmp.max_row-1,'csv_rows':len(export)},indent=2))
